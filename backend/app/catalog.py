"""Staff catalogue maintenance and review. Import previews never mutate stock."""
import csv
import hashlib
import io
import json
import zipfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from urllib.parse import quote
from . import db, service

HEADERS = ['操作', '商品名称', '型号', '分类', '商品说明', '规格名称', '颜色', '尺码', '售价（元）', '库存']
MAX_ROWS = 500
SCHEMA = '''
CREATE TABLE IF NOT EXISTS catalog_imports(id TEXT PRIMARY KEY,user_id TEXT NOT NULL,fingerprint TEXT NOT NULL,data TEXT NOT NULL,created_at TEXT NOT NULL,confirmed_at TEXT,UNIQUE(user_id,fingerprint));
CREATE TABLE IF NOT EXISTS answer_drafts(id TEXT PRIMARY KEY,product_id TEXT NOT NULL,status TEXT NOT NULL,version INTEGER NOT NULL DEFAULT 1,data TEXT NOT NULL,created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
'''

def init_schema(c):
    c.executescript(SCHEMA)


def product_path(pid, sku_id=None):
    path = '/products/' + quote(str(pid), safe='')
    return path + ('?sku=' + quote(str(sku_id), safe='') if sku_id else '')


def decorate_product(p):
    p['purchase_path'] = product_path(p['id'])
    p['catalog_version'] = p.get('catalog_version', 1)
    for sku in p.get('skus', []):
        sku['purchase_path'] = product_path(p['id'], sku['id'])
    return p


def _text(data, key, label, required=True, maximum=200):
    value = data.get(key, '')
    if value is None: value = ''
    if not isinstance(value, str): service.fail(422, f'{label}请填写文字')
    value = value.strip()
    if required and not value: service.fail(422, f'请填写{label}')
    if len(value) > maximum: service.fail(422, f'{label}不能超过 {maximum} 个字')
    return value


def _integer(value, label, low=0, high=1_000_000):
    if isinstance(value, bool): service.fail(422, f'{label}请填写整数')
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value() or not low <= number <= high: raise ValueError()
    except (InvalidOperation, ValueError): service.fail(422, f'{label}请填写 {low} 至 {high} 的整数')
    return int(number)


def _price(value):
    try:
        price = Decimal(str(value))
        if not price.is_finite() or not Decimal('0.01') <= price <= Decimal('1000000') or price * 100 != (price * 100).to_integral_value(): raise ValueError()
    except (InvalidOperation, ValueError): service.fail(422, '售价请填写大于零且最多两位小数的金额')
    return int(price * 100)


def _sku_data(data):
    return {'label': _text(data, 'label', '规格名称', maximum=100),
            'color': _text(data, 'color', '颜色', False, 50), 'size': _text(data, 'size', '尺码', False, 50),
            'price_cents': _price(data.get('price')), 'stock': _integer(data.get('stock'), '库存')}


def _base_data(data):
    values = {key: _text(data, key, label, key != 'model', maximum) for key, label, maximum in [
        ('name', '商品名称', 100), ('model', '型号', 80), ('category', '分类', 40), ('description', '商品说明', 5000)]}
    return values


def _natural_exists(c, name, model, except_id=None):
    for row in c.execute('SELECT id,data FROM products'):
        p = json.loads(row['data'])
        if row['id'] != except_id and p['name'].casefold() == name.casefold() and p['model'].casefold() == model.casefold():
            service.fail(409, '相同名称和型号的商品已存在，请修改已有商品')


def _insert(c, user, base, skus):
    _natural_exists(c, base['name'], base['model'])
    pid = db.uid('product')
    product = {**base, 'id': pid, 'image': '/assets/catalog-placeholder.svg', 'aliases': list(filter(None, [base['name'], base['model']])),
               'price_cents': min(s['price_cents'] for s in skus), 'specs': {}, 'faq': [], 'size_chart': None,
               'package_contents': ['商品本体 × 1'], 'package_items': [{'id': pid + '-main', 'name': '商品本体', 'kind': 'main', 'quantity': 1}],
               'source_url': '', 'retrieved_at': db.now(), 'catalog_version': 1}
    c.execute('INSERT INTO products VALUES (?,?)', (pid, service.js(product)))
    for sku in skus:
        sid = db.uid('sku')
        c.execute('INSERT INTO skus(id,product_id,label,color,size,price_cents,stock) VALUES (?,?,?,?,?,?,?)',
                  (sid, pid, sku['label'], sku['color'], sku['size'], sku['price_cents'], sku['stock']))
        c.execute('INSERT INTO inventory_movements VALUES (?,?,?,?,?,?)', (db.uid('stock'), sid, sku['stock'], '新增商品库存', pid, db.now()))
    service.audit(c, user, '新增商品', pid, None, product)
    return service._product(c, pid)


def create_product(user, data):
    service.require_staff(user)
    request_id = _text(data, 'request_id', '提交凭据')
    base = _base_data(data)
    if not isinstance(data.get('skus'), list) or not 1 <= len(data['skus']) <= 100: service.fail(422, '请添加 1 至 100 个规格')
    skus = [_sku_data(s) for s in data['skus'] if isinstance(s, dict)]
    if len(skus) != len(data['skus']): service.fail(422, '请检查商品规格')
    if len({s['label'].casefold() for s in skus}) != len(skus): service.fail(422, '同一商品的规格名称不能重复')
    with db.transaction() as c:
        previous = service.memo(c, user, 'catalog-create', request_id)
        if previous: return previous
        return service.remember(c, user, 'catalog-create', request_id, _insert(c, user, base, skus))


def edit_product(user, pid, data):
    service.require_staff(user)
    base = _base_data(data)
    specs = data.get('specs', {})
    if not isinstance(specs, dict) or len(specs) > 40 or any(not isinstance(k,str) or not k.strip() or not isinstance(v,str) or len(v)>1000 for k,v in specs.items()):
        service.fail(422, '请用文字填写商品参数，最多 40 项')
    with db.transaction() as c:
        old = service._product(c, pid)
        if data.get('expected_version') != old.get('catalog_version', 1): service.fail(409, '这款商品已被修改，请刷新后再保存')
        _natural_exists(c, base['name'], base['model'], pid)
        new = {k:v for k,v in old.items() if k not in ('skus','parts','purchase_path')}
        new.update(base, specs=specs, catalog_version=old.get('catalog_version',1)+1, retrieved_at=db.now())
        c.execute('UPDATE products SET data=? WHERE id=?', (service.js(new),pid))
        service.audit(c,user,'更新商品资料',pid,old,new)
        return service._product(c,pid)


def add_sku(user, pid, data):
    service.require_staff(user)
    sku = _sku_data(data)
    request_id = _text(data, 'request_id', '提交凭据')
    with db.transaction() as c:
        previous = service.memo(c,user,'catalog-sku-'+pid,request_id)
        if previous: return previous
        p = service._product(c,pid)
        if any(s['label'].casefold()==sku['label'].casefold() for s in p['skus']): service.fail(409,'这个规格已经存在，请修改已有规格')
        sid=db.uid('sku')
        c.execute('INSERT INTO skus(id,product_id,label,color,size,price_cents,stock) VALUES (?,?,?,?,?,?,?)',(sid,pid,sku['label'],sku['color'],sku['size'],sku['price_cents'],sku['stock']))
        c.execute('INSERT INTO inventory_movements VALUES (?,?,?,?,?,?)',(db.uid('stock'),sid,sku['stock'],'新增规格库存',pid,db.now()))
        service.audit(c,user,'新增商品规格',sid,None,sku)
        return service.remember(c,user,'catalog-sku-'+pid,request_id,service._product(c,pid))


def _check_xlsx_bounds(sheet):
    """Check actual cells before a bounded read; formatting alone is harmless.

    The read-only worksheet source is scanned as sparse XML so a formatted cell
    at Excel's last row/column cannot expand into a huge rectangular iteration.
    Uses the private read-only reader hooks from the pinned openpyxl 3.1 series.
    """
    from xml.etree.ElementTree import iterparse
    from openpyxl.utils.cell import coordinate_to_tuple
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with sheet._get_source() as source:
        for _, cell in iterparse(source, events=('end',)):
            if cell.tag == ns + 'row':
                cell.clear()
                continue
            if cell.tag != ns + 'c': continue
            row, column = coordinate_to_tuple(cell.attrib['r'])
            if row <= MAX_ROWS + 1 and column <= len(HEADERS):
                cell.clear()
                continue
            formula = cell.find(ns + 'f')
            values = [part.text for part in cell.iter() if part.tag in (ns + 'v', ns + 't')]
            if cell.attrib.get('t') == 's' and values and values[0] is not None:
                values = [sheet._shared_strings[int(values[0])]]
            populated = formula is not None or any(value is not None and str(value).strip() for value in values)
            cell.clear()
            if not populated: continue
            if column > len(HEADERS):
                service.fail(422, f'第 {row} 行第 {column} 列在模板之外还有内容，请移回对应列或删除多余内容；本次未导入任何行')
            service.fail(422, f'第 {row} 行超出模板范围，一次最多导入 {MAX_ROWS} 行商品规格，请分批上传；本次未导入任何行')


def _read_file(filename, content):
    if len(content)>2*1024*1024: service.fail(413,'表格不能超过 2 MB，请分批导入')
    if filename.lower().endswith('.csv'):
        try: text=content.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: text=content.decode('gb18030')
            except UnicodeDecodeError: service.fail(422,'表格文字无法识别，请另存为中文 CSV 或 Excel 表格')
        rows=list(csv.reader(io.StringIO(text)))
    elif filename.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist())>20*1024*1024: service.fail(413,'表格内容过大，请删除空白区域后重试')
            from openpyxl import load_workbook
            workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
            try:
                sheet=workbook.active
                _check_xlsx_bounds(sheet)
                rows=list(sheet.iter_rows(max_row=MAX_ROWS+2,max_col=len(HEADERS)+1,values_only=True))
            finally: workbook.close()
        except service.HTTPException: raise
        except Exception: service.fail(422,'无法读取此表格，请使用下载的模板另存为 Excel 表格')
    else: service.fail(422,'请选择 Excel 表格或 CSV 表格')
    while rows and not any(v is not None and str(v).strip() for v in rows[-1]): rows.pop()
    if len(rows)<2: service.fail(422,'表格里没有商品，请先填写商品再上传')
    if len(rows)>MAX_ROWS+1: service.fail(422,f'一次最多导入 {MAX_ROWS} 行商品规格，请分批上传')
    headers=[str(v or '').strip() for v in rows[0]]
    while headers and not headers[-1]: headers.pop()
    if headers!=HEADERS: service.fail(422,'表头与模板不一致，请下载模板并保留第一行列名')
    return [{**{h:row[i] if i<len(row) and row[i] is not None else '' for i,h in enumerate(HEADERS)}, '_extra_columns': any(v is not None and str(v).strip() for v in row[len(HEADERS):])} for row in rows[1:]]


def _fingerprint(value):
    return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True).encode()).hexdigest()


def preview_import(user,filename,content):
    service.require_staff(user)
    rows=_read_file(filename,content)
    fingerprint=_fingerprint(rows)
    with db.transaction() as c:
        cached=c.execute('SELECT * FROM catalog_imports WHERE user_id=? AND fingerprint=?',(user['id'],fingerprint)).fetchone()
        if cached and cached['confirmed_at']:
            previous=json.loads(cached['data']); previous.update(id=cached['id'],already_confirmed=True)
            return {k:v for k,v in previous.items() if k!='operations'}
        products=service._products(c)
        result=[]; ops=[]; seen=set(); new_groups={}
        for i,row in enumerate(rows,2):
            errors=[]; operation=None
            try:
                if row.get('_extra_columns'): service.fail(422,'模板列之外还有内容，请移回对应列或删除多余内容')
                if any(isinstance(v,str) and v.lstrip().startswith('=') for v in row.values()): service.fail(422,'请填写具体文字或数字，不要使用公式')
                mode=str(row['操作']).strip()
                if mode not in ('新增','更新'): service.fail(422,'操作请填写“新增”或“更新”')
                name=str(row['商品名称']).strip(); model=str(row['型号']).strip()
                if not name: service.fail(422,'请填写商品名称')
                sku=_sku_data({'label':str(row['规格名称']),'color':str(row['颜色']),'size':str(row['尺码']),'price':row['售价（元）'],'stock':row['库存']})
                key=(name.casefold(),model.casefold(),sku['label'].casefold())
                if key in seen: service.fail(422,'与表格前面的商品规格重复，请只保留一行')
                seen.add(key)
                matches=[p for p in products if p['name'].casefold()==name.casefold() and p['model'].casefold()==model.casefold()]
                if mode=='新增':
                    if matches: service.fail(422,'商品已经存在，修改价格库存请把操作改为“更新”')
                    base=_base_data({'name':name,'model':model,'category':str(row['分类']),'description':str(row['商品说明'])})
                    group=(name.casefold(),model.casefold())
                    if group in new_groups and new_groups[group]!=base: service.fail(422,'同一商品的分类或说明不一致，请统一后导入')
                    new_groups[group]=base
                    operation={'mode':'create','base':base,'sku':sku}
                else:
                    if len(matches)!=1: service.fail(422,'没有唯一对应的商品，请核对商品名称和型号')
                    product=matches[0]
                    variants=[s for s in product['skus'] if s['label']==sku['label']]
                    if len(variants)!=1: service.fail(422,'没有对应规格，请先新增规格或核对规格名称')
                    before=variants[0]
                    if sku['color'] and sku['color']!=before['color'] or sku['size'] and sku['size']!=before['size']: service.fail(422,'颜色或尺码与现有规格不一致，请核对后重试')
                    operation={'mode':'update','pid':product['id'],'sid':before['id'],'expected_version':before['version'],'sku':sku}
                ops.append(operation)
            except service.HTTPException as e: errors.append(str(e.detail))
            result.append({'row_number':i,'values':{k:str(v) for k,v in row.items() if k in HEADERS},'errors':errors})
        preview={'rows':result,'valid_count':sum(not r['errors'] for r in result),'error_count':sum(bool(r['errors']) for r in result),'already_confirmed':False,'operations':ops}
        iid=cached['id'] if cached else db.uid('import')
        if cached: c.execute('UPDATE catalog_imports SET data=?,created_at=? WHERE id=?',(service.js(preview),db.now(),iid))
        else: c.execute('INSERT INTO catalog_imports(id,user_id,fingerprint,data,created_at) VALUES (?,?,?,?,?)',(iid,user['id'],fingerprint,service.js(preview),db.now()))
        return {k:v for k,v in dict(preview,id=iid).items() if k!='operations'}


def confirm_import(user,iid):
    service.require_staff(user)
    with db.transaction() as c:
        row=service.one(c,'SELECT * FROM catalog_imports WHERE id=? AND user_id=?',(iid,user['id']),'导入预览不存在，请重新上传表格')
        preview=json.loads(row['data'])
        if row['confirmed_at']: return {'count':preview['valid_count'],'already_confirmed':True}
        if preview['error_count']: service.fail(422,'表格还有错误，请修正后重新上传，尚未写入任何商品')
        if datetime.now(timezone.utc)-datetime.fromisoformat(row['created_at'])>timedelta(hours=2): service.fail(409,'预览已过期，请重新上传核对最新资料')
        groups={}
        for op in preview['operations']:
            if op['mode']=='create':
                key=(op['base']['name'],op['base']['model']); groups.setdefault(key,{'base':op['base'],'skus':[]})['skus'].append(op['sku'])
            else:
                old=service._sku(c,op['sid'])
                if old['version']!=op['expected_version']: service.fail(409,'预览后商品价格或库存发生变化，请重新上传核对；本次未写入任何行')
                sku=op['sku']
                c.execute('UPDATE skus SET price_cents=?,stock=?,version=version+1 WHERE id=?',(sku['price_cents'],sku['stock'],op['sid']))
                c.execute('INSERT INTO inventory_movements VALUES (?,?,?,?,?,?)',(db.uid('stock'),op['sid'],sku['stock']-old['stock'],'表格更新库存',iid,db.now()))
                service.audit(c,user,'表格更新商品规格',op['sid'],old,sku)
        for group in groups.values(): _insert(c,user,group['base'],group['skus'])
        c.execute('UPDATE catalog_imports SET confirmed_at=? WHERE id=?',(db.now(),iid))
        service.audit(c,user,'确认商品表格导入',iid,None,{'count':preview['valid_count']})
        return {'count':preview['valid_count'],'already_confirmed':False}


def template(user,format='xlsx',current=False):
    service.require_staff(user)
    rows=[HEADERS]
    if current:
        for p in service.products():
            for s in p['skus']: rows.append(['更新',p['name'],p['model'],p['category'],p['description'],s['label'],s['color'],s['size'],format_price(s['price_cents']),s['stock']])
    else: rows.append(['新增','示例收纳盒','日常款','家居','可叠放，便于整理桌面','透明 / 小号','透明','小号','19.90',20])
    if format=='csv':
        out=io.StringIO(); writer=csv.writer(out)
        writer.writerows([['\''+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for v in r] for r in rows])
        return out.getvalue().encode('utf-8-sig')
    from openpyxl import Workbook
    book=Workbook(); sheet=book.active; sheet.title='商品资料'
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'
    sheet.freeze_panes='A2'
    from openpyxl.styles import Font, PatternFill
    for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='28604E')
    for col in 'ABCDEFGHIJ': sheet.column_dimensions[col].width=22 if col!='E' else 45
    output=io.BytesIO(); book.save(output); return output.getvalue()


def format_price(cents): return str(Decimal(cents)/100)


def _faq(value):
    if not isinstance(value,list) or len(value)>30: service.fail(422,'标准答案最多 30 条')
    result=[]
    for item in value:
        if not isinstance(item,dict): service.fail(422,'请完整填写问题和答案')
        result.append({'q':_text(item,'q','问题',maximum=300),'a':_text(item,'a','答案',maximum=3000)})
    if len({v['q'] for v in result})!=len(result): service.fail(422,'相同问题请只保留一个答案')
    return result


def answer_drafts(user):
    service.require_staff(user)
    with db.connect() as c:
        return [dict(json.loads(row['data']),id=row['id'],product_id=row['product_id'],status=row['status'],version=row['version'],created_at=row['created_at']) for row in c.execute('SELECT * FROM answer_drafts ORDER BY created_at DESC LIMIT 100')]


def submit_answers(user,pid,data):
    service.require_staff(user); faq=_faq(data.get('faq')); req=_text(data,'request_id','提交凭据')
    with db.transaction() as c:
        previous=service.memo(c,user,'answer-submit',req)
        if previous: return previous
        p=service._product(c,pid)
        if p['faq']==faq: service.fail(422,'答案没有变化，无需重复提交')
        if data.get('expected_version')!=p.get('catalog_version',1): service.fail(409,'商品资料已更新，请刷新后提交')
        draft={'product_name':p['name'],'before_faq':p['faq'],'faq':faq,'submitted_by':user['name'],'review_note':''}
        did=db.uid('answer'); stamp=db.now()
        c.execute('INSERT INTO answer_drafts(id,product_id,status,data,created_at,updated_at) VALUES (?,?,?,?,?,?)',(did,pid,'pending',service.js(draft),stamp,stamp))
        service.audit(c,user,'提交标准答案审核',pid,None,draft)
        return service.remember(c,user,'answer-submit',req,{'id':did,'status':'pending'})


def review_answers(user,did,data):
    service.require_staff(user)
    decision=data.get('decision')
    if decision not in ('approve','reject'): service.fail(422,'请选择通过或退回')
    with db.transaction() as c:
        row=service.one(c,'SELECT * FROM answer_drafts WHERE id=?',(did,),'没有找到这份答案审核')
        if row['status']!='pending': service.fail(409,'这份答案已处理，请刷新列表')
        if data.get('expected_version')!=row['version']: service.fail(409,'审核内容已经变化，请刷新后重试')
        draft=json.loads(row['data']); p=service._product(c,row['product_id'])
        if decision=='approve':
            if p['faq']!=draft['before_faq']: service.fail(409,'现行答案已经改变，请重新提交核对，避免覆盖新答案')
            new={k:v for k,v in p.items() if k not in ('skus','parts','purchase_path')}
            new.update(faq=draft['faq'],catalog_version=p.get('catalog_version',1)+1,retrieved_at=db.now())
            c.execute('UPDATE products SET data=? WHERE id=?',(service.js(new),p['id']))
        draft.update(review_note=_text(data,'note','审核说明',False,1000),reviewed_by=user['name'],reviewed_at=db.now())
        status='approved' if decision=='approve' else 'rejected'
        c.execute('UPDATE answer_drafts SET status=?,version=version+1,data=?,updated_at=? WHERE id=?',(status,service.js(draft),db.now(),did))
        service.audit(c,user,'通过标准答案' if decision=='approve' else '退回标准答案',p['id'],draft['before_faq'],draft)
        return {'id':did,'status':status}
