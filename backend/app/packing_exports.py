"""Chinese review documents generated from one current packing revision."""
import io
import json
from urllib.parse import quote
from xml.sax.saxutils import escape
from fastapi import APIRouter, Depends
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from . import db, packing, packing_pdf, packing_workflow, service


def _text(value):
    if value is None or value == '':
        return '未提供'
    if isinstance(value, bool):
        return '是' if value else '否'
    if isinstance(value, (list, tuple)):
        return '、'.join(_text(v) for v in value)
    return str(value)


def _source_rows(task):
    snapshot = task.get('intake_snapshot') or {}
    fields = snapshot.get('fields', [])
    sources = {s.get('id'): s for s in snapshot.get('sources', [])}
    rows = []
    for field in fields:
        refs = []
        for cite in field.get('citations', []):
            source = sources.get(cite.get('source_id'), {})
            chunk = next((v for v in source.get('chunks', []) if v.get('id') == cite.get('chunk_id')), {})
            loc = chunk.get('locator', {})
            if cite.get('image_id'):
                page = next((v for v in source.get('pages', []) if v.get('image_id') == cite['image_id']), {})
                loc = {'page': page.get('page')}
            position = ''.join([f'第{loc["page"]}页' if loc.get('page') else '',
                                f'{loc["sheet"]}工作表 ' if loc.get('sheet') else '',
                                f'{loc["cell"]}单元格 ' if loc.get('cell') else '',
                                f'第{loc["row"]}行 ' if loc.get('row') else '',
                                f'第{loc["paragraph"]}段 ' if loc.get('paragraph') else '',
                                f'第{loc["table"]}个表格 ' if loc.get('table') else ''])
            refs.append(f'《{source.get("name", cite.get("source_name", "原始资料"))}》{position}：{cite.get("quote", "请核对原图")}')
        reason = field.get('confirmation', {}).get('reason') if field.get('confirmation') else None
        if reason:
            refs.append('人工核对说明：' + reason)
        candidates = field.get('candidates', [])
        if field.get('name') != 'required_fields' and len({str(v.get('value')) for v in candidates}) > 1:
            refs.append('曾出现不同内容：' + '；'.join(dict.fromkeys(_text(v.get('value')) for v in candidates)))
        for rejected in field.get('rejected_candidates', []):
            refs.append('未采用的候选：' + _text(rejected.get('value')) + '；原因：' + str(rejected.get('reason', rejected.get('error', '未通过资料核对'))))
        value = field.get('value')
        if field.get('name') == 'required_fields' and isinstance(value, list):
            from .packing_sources import TITLES
            value = [TITLES.get(name.split('.')[-1], '其他必填资料') for name in value]
        rows.append([field.get('title', '订单资料'), _text(value), '\n'.join(refs) or '人工填写；请核对录入依据',
                     '已人工核对' if field.get('confirmed') else '待核对'])
    if not rows:
        for title, key in [('发货单位', 'supplier'), ('收货仓库', 'destination'), ('客户订单号', 'platform_order'), ('计划发货日', 'ship_date')]:
            rows.append([title, _text(task.get(key)), '员工手工录入', '最终审核已通过' if task.get('approval') else '待最终审核'])
        for i, line in enumerate(task['lines'], 1):
            rows.extend([[f'第{i}行商品／规格', line['product_name']+'／'+line['specification'], '店内商品资料及员工选择', '已选择'],
                         [f'第{i}行发货数量', str(line['quantity']), '员工手工录入', '待最终审核' if not task.get('approval') else '最终审核已通过'],
                         [f'第{i}行每箱数量', str(line['per_carton']), '员工手工录入', '待最终审核' if not task.get('approval') else '最终审核已通过']])
    for i, line in enumerate(task['lines'], 1):
        full, tail = divmod(line['quantity'], line['per_carton'])
        rows.append([f'第{i}行分箱依据', f'{line["quantity"]}件，共{full + bool(tail)}箱',
                     f'每箱{line["per_carton"]}件 × {full}整箱' + (f' ＋ 尾箱{tail}件' if tail else ''), '按已核对数量计算'])
    for material in task.get('materials', []):
        count = task['lines'][material['line_index']]['quantity']
        rows.append([material['name']+'备料依据', f'{material["planned_quantity"]}件',
                     f'成品{count}件 × 每套{material["per_item"]}件配件', '按包装资料计算'])
    return rows


def _change_record(change):
    shown = packing_workflow.display_change(change)
    notes = []
    if shown.get('resolution'):
        notes.append('人工处理记录：' + shown['resolution'])
    for item in shown.get('differences', []):
        if item.get('label') == '原资料核对版本':
            notes.append('来源版本申请对比：' + _text(item.get('before')) + ' → ' + _text(item.get('after')))
    if shown.get('impact'):
        notes.append('影响说明：' + '；'.join(shown['impact']))
    return '\n'.join(notes)


def _tables(task, changes):
    return {
        '订单核对': [['项目', '当前内容'], ['用途', '客户PO核对整理件，不代替客户原始采购单'], ['装箱任务', task['number']],
                   ['资料版本', task['revision']], ['当前进度', task['stage_label']], ['发货单位', task['supplier']], ['收货仓库', task['destination']],
                   ['客户订单号', task['platform_order']], ['计划发货日', task['ship_date']], ['发货数量（件）', task['total_quantity']], ['箱数', task['total_boxes']]],
        '商品明细': [['内部商品编号', '商品', '规格', '发货数量（件）', '每箱数量（件）', '批次', '商品条码', '客户商品编码']]+
                    [[l['sku_id'], l['product_name'], l['specification'], l['quantity'], l['per_carton'], l['batch'], l['product_barcode'], l.get('external_code', '')] for l in task['lines']],
        '逐箱清单': [['箱号', '商品', '规格', '计划数量', '实装数量', '净重（公斤）', '毛重（公斤）', '批次', '序列号', '当前核验号']]+
                    [[b['number'], task['lines'][b['line_index']]['product_name'], task['lines'][b['line_index']]['specification'],
                      b['planned_quantity'], b['actual_quantity'], b['net_weight'], b['gross_weight'], task['lines'][b['line_index']]['batch'],
                      '、'.join(b['serial_numbers']), b['code'] or '草稿无有效核验号'] for b in task['boxes']],
        '配件备料': [['商品', '配件', '每套用量', '应备数量', '实备数量']]+
                    [[task['lines'][m['line_index']]['product_name'], m['name'], m['per_item'], m['planned_quantity'], m['actual_quantity']] for m in task.get('materials', [])],
        '资料出处': [['核对项目', '当前内容', '来源及原文', '核对结果']]+_source_rows(task),
        '变更处理': [['时间', '变更原因', '通知出处', '处理进度', '处理记录']]+
                    [[v['created_at'], v['reason'], v.get('source_note', ''), {'pending':'待确认', 'confirmed':'已确认重建', 'requires_followup':'出库后待人工处理', 'resolved':'已登记人工处理', 'withdrawn':'已撤销', 'superseded':'已有其他变更生效'}.get(v['status'], '待核对'),
                      _change_record(v)] for v in changes],
    }


def _xlsx(task, changes):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in _tables(task, changes).items():
        ws = wb.create_sheet(name)
        ws.append(['演示资料；仓库要求待核对。导出为当前版本记录，不表示已接入ERP或WMS。'])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(rows[0])))
        ws.row_dimensions[1].height = 34
        for row in rows:
            ws.append(row)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Values from documents never become workbook formulas.
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='等线', size=11)
        for cell in ws[2]:
            cell.fill = PatternFill('solid', fgColor='214C40')
            cell.font = Font(name='等线', size=11, color='FFFFFF', bold=True)
        for column in range(1, len(rows[0])+1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(column)].width = 25 if name != '资料出处' else [23, 28, 78, 20][column-1]
        ws.freeze_panes = 'A3'
        if len(rows) > 1:
            ws.auto_filter.ref = f'A2:{ws.cell(ws.max_row, len(rows[0])).coordinate}'
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = '1:2'
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _pdf(task, changes, kind):
    packing_pdf._font()
    output = io.BytesIO()
    title = {'po': '客户PO核对整理件', 'manifest': '逐箱装箱与配件清单', 'sources': '资料来源与异常处理报告'}[kind]
    normal = ParagraphStyle('body', fontName=packing_pdf.FONT, fontSize=9, leading=14, wordWrap='CJK', spaceAfter=4)
    heading = ParagraphStyle('title', parent=normal, fontSize=18, leading=25, spaceAfter=12)
    def p(value):
        return Paragraph(escape(_text(value)).replace('\n', '<br/>'), normal)
    story = [Paragraph(title, heading), p(f'{task["number"]} · 第{task["revision"]}版 · {task["stage_label"]}'),
             p('演示资料，仓库要求待核对。客户PO原件保持不变；本报告为核对整理记录。'), Spacer(1, 5*mm)]
    tables = _tables(task, changes)
    names = ['订单核对', '商品明细'] if kind == 'po' else ['逐箱清单', '配件备料'] if kind == 'manifest' else ['资料出处', '变更处理']
    for name in names:
        story += [Paragraph(name, ParagraphStyle('section', parent=normal, fontSize=12, leading=18, spaceBefore=10)), Spacer(1, 2*mm)]
        rows = tables[name]
        if name == '逐箱清单':
            rows = [[*r[:5], f'净重：{_text(r[5])}\n毛重：{_text(r[6])}\n批次：{_text(r[7])}\n序列号：{_text(r[8])}\n核验号：{_text(r[9])}'] for r in rows]
            rows[0][-1] = '重量与追溯资料'
            widths = [12, 27, 23, 16, 16, 80]
        elif name == '资料出处':
            widths = [28, 30, 91, 25]
        elif name == '订单核对':
            widths = [40, 134]
        elif name == '商品明细':
            widths = [24, 36, 24, 20, 20, 22, 28]
        else:
            widths = [174/len(rows[0])] * len(rows[0])
        if len(rows) == 1:
            story.append(p('暂无相关记录。'))
            continue
        # splitInRow permits a long source quote or serial list to continue without truncation.
        table = Table([[p(v) for v in row] for row in rows], colWidths=[v*mm for v in widths], repeatRows=1, splitByRow=1, splitInRow=1)
        table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E8EFEA')),
                                  ('GRID', (0,0), (-1,-1), .35, colors.HexColor('#CFD8D2')),
                                  ('VALIGN',(0,0),(-1,-1),'TOP'), ('LEFTPADDING',(0,0),(-1,-1),6),
                                  ('RIGHTPADDING',(0,0),(-1,-1),6), ('TOPPADDING',(0,0),(-1,-1),7), ('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        story.append(table)
    def footer(canvas, document):
        canvas.setFont(packing_pdf.FONT, 8)
        canvas.drawString(18*mm, 12*mm, f'{task["number"]} · 第{task["revision"]}版 · 演示资料')
        canvas.drawRightString(A4[0]-18*mm, 12*mm, f'第{document.page}页')
    SimpleDocTemplate(output, pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=22*mm,
                      title=title, author='日常好物箱唛工作台').build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export(user, identity, kind):
    packing._staff(user)
    if kind not in ('po', 'manifest', 'sources', 'xlsx'):
        service.fail(422, '请选择订单核对、装箱清单、来源报告或Excel')
    with db.connect() as c:
        task = packing._get(c, identity)
        if task.get('approval'):
            packing_workflow.check_approval(c, task)
        changes = [json.loads(r[0]) for r in c.execute('SELECT data FROM packing_changes WHERE task_id=? ORDER BY rowid', (identity,))]
    labels = {'po': '客户PO核对整理件', 'manifest': '逐箱装箱清单', 'sources': '资料来源与异常报告', 'xlsx': '订单与装箱资料'}
    content = _xlsx(task, changes) if kind == 'xlsx' else _pdf(task, changes, kind)
    suffix = 'xlsx' if kind == 'xlsx' else 'pdf'
    filename = f'{task["number"]}-第{task["revision"]}版-{labels[kind]}.{suffix}'
    mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if kind == 'xlsx' else 'application/pdf'
    return content, filename, mime


def build_router(user_dependency):
    router = APIRouter(prefix='/api/packing')
    @router.get('/tasks/{identity}/exports/{kind}')
    def download(identity: str, kind: str, u=Depends(user_dependency)):
        content, filename, mime = export(u, identity, kind)
        return Response(content, media_type=mime, headers={'Content-Disposition': f"attachment; filename*=UTF-8''{quote(filename)}", 'Cache-Control':'private, no-store'})
    return router
