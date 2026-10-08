"""Bounded, non-executing readers for employee packing source documents.

Only this module resolves persisted storage names. User input never supplies a
filesystem path. The original bytes are content addressed and never overwritten.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, UnidentifiedImageError

from . import db, service
from .config import settings

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_EXPANDED_BYTES = 40 * 1024 * 1024
MAX_TEXT = 100000
MAX_CHUNKS = 6000
MAX_PAGES = 20
MAX_PIXELS = 24000000
EXTENSIONS = {'.xlsx': 'xlsx', '.csv': 'csv', '.docx': 'docx', '.pdf': 'pdf',
              '.png': 'image', '.jpg': 'image', '.jpeg': 'image', '.webp': 'image'}
MEDIA = {'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
         'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         'pdf': 'application/pdf', 'csv': 'text/csv', 'text': 'text/plain'}


def storage_path(name):
    if not isinstance(name, str) or not re.fullmatch(r'(?:blobs|pages)/[a-zA-Z0-9_.-]+', name):
        service.fail(422, '资料存储标识不正确')
    root = (settings.upload_dir / 'packing').resolve()
    path = (root / name).resolve()
    if not path.is_relative_to(root):
        service.fail(422, '资料存储标识不正确')
    return path


def _write_once(name, content):
    path = storage_path(name)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with path.open('xb') as target:
            target.write(content)
    except FileExistsError:
        if not path.is_file() or hashlib.sha256(path.read_bytes()).digest() != hashlib.sha256(content).digest():
            service.fail(409, '原资料存储核对失败，请联系维护人员')
    return name


def _chunk(text, locator, **extra):
    return {'id': db.uid('chunk'), 'text': str(text), 'locator': locator, **extra}


def _bounded(chunks):
    if len(chunks) > MAX_CHUNKS or sum(len(chunk['text']) for chunk in chunks) > MAX_TEXT:
        service.fail(422, '资料文字或单元格过多，请拆分资料后上传')
    return chunks


def _zip(content, required):
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
        items = archive.infolist()
        if len(items) > 3000 or sum(item.file_size for item in items) > MAX_EXPANDED_BYTES:
            service.fail(422, '文件解压后过大，请拆分资料后上传')
        for item in items:
            name = item.filename.replace('\\', '/')
            if name.startswith('/') or '..' in name.split('/') or item.flag_bits & 1:
                service.fail(422, '文件包含不安全或加密内容，请重新保存后上传')
            if item.file_size > 1024 * 1024 and item.file_size > max(1, item.compress_size) * 200:
                service.fail(422, '文件压缩比例异常，请重新保存后上传')
            if re.search(r'(?i)(vbaproject|activex|externalLinks|embeddings|\.exe$|\.js$)', name):
                service.fail(422, '资料含宏、外部链接或嵌入程序，请移除后上传')
        if required not in archive.namelist():
            service.fail(422, '文件内容与扩展名不符，请上传有效的资料文件')
        return archive
    except (zipfile.BadZipFile, OSError, ValueError):
        service.fail(422, '无法读取文件，请重新保存为有效的资料格式')


def _xml(content):
    if len(content) > MAX_EXPANDED_BYTES or re.search(br'<!\s*(?:DOCTYPE|ENTITY)', content, re.I):
        service.fail(422, '文件包含不安全的文档结构，请重新保存后上传')
    try:
        return ET.fromstring(content)
    except ET.ParseError:
        service.fail(422, '文档结构损坏，请重新保存后上传')


def _cell_text(cell):
    value = cell.value
    if value is None:
        return ''
    if cell.data_type == 'f':
        service.fail(422, f'{cell.coordinate} 包含公式，请复制为实际值后上传；系统不会执行公式')
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, (int, float)):
        fmt = str(cell.number_format or '')
        if re.fullmatch(r'0{2,30}', fmt) and float(value).is_integer() and value >= 0:
            return str(int(value)).zfill(len(fmt))
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


def read_xlsx(content):
    from openpyxl import load_workbook
    with _zip(content, 'xl/workbook.xml') as archive:
        for name in archive.namelist():
            if name.endswith('.rels'):
                rels = _xml(archive.read(name))
                if any(item.get('TargetMode') == 'External' for item in rels):
                    service.fail(422, '表格含外部链接，请移除后上传')
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception:
        service.fail(422, '无法读取表格，请保存为 xlsx 格式后上传')
    chunks, notes = [], []
    try:
        if len(book.worksheets) > 20:
            service.fail(422, '工作表过多，请拆分资料后上传')
        for sheet in book.worksheets:
            if (sheet.max_row or 0) > 2000 or (sheet.max_column or 0) > 60:
                service.fail(422, '表格最多读取 2000 行、60 列，请拆分资料后上传')
            headers = []
            for row_index, cells in enumerate(sheet.iter_rows(), 1):
                if row_index > 2000 or len(cells) > 60:
                    service.fail(422, '表格最多读取 2000 行、60 列，请拆分资料后上传')
                pairs, values = [], []
                for column, cell in enumerate(cells):
                    value = _cell_text(cell)
                    values.append(value)
                    if value:
                        chunks.append(_chunk(value, {'sheet': sheet.title, 'cell': cell.coordinate, 'row': row_index, 'column': column + 1},
                                             raw_value=str(cell.value), number_format=str(cell.number_format)))
                        header = headers[column] if column < len(headers) and headers[column] else cell.coordinate
                        pairs.append(header + '=' + value)
                        if isinstance(cell.value, (int, float)) and re.search(r'编码|条码|序列|订单号|SKU|code', header, re.I):
                            if not re.fullmatch(r'0{2,30}', str(cell.number_format)):
                                notes.append(f'{sheet.title}!{cell.coordinate} 的编号为数值，请核对前导零及长编号精度')
                if row_index == 1:
                    headers = values
                if pairs:
                    chunks.append(_chunk(' | '.join(pairs), {'sheet': sheet.title, 'row': row_index}, kind='table_row'))
                _bounded(chunks)
        return chunks, [], list(dict.fromkeys(notes))
    finally:
        book.close()


def read_csv(content):
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            text = content.decode('gb18030')
        except UnicodeDecodeError:
            service.fail(422, 'CSV 编码无法读取，请保存为 UTF-8 后上传')
    if len(text) > MAX_TEXT or '\x00' in text:
        service.fail(422, 'CSV 文字过多或包含无效内容，请拆分后上传')
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    chunks, headers = [], []
    try:
        for index, row in enumerate(csv.reader(io.StringIO(text, newline=''), dialect), 1):
            if index > 2000 or len(row) > 60:
                service.fail(422, 'CSV 最多读取 2000 行、60 列，请拆分后上传')
            if any(value.lstrip().startswith(('=', '+', '-', '@')) and re.search(r'[=()]|^[+@]', value.lstrip()) for value in row):
                service.fail(422, f'CSV 第 {index} 行含公式样式内容，请复制为实际值后上传')
            if index == 1:
                headers = row
            pairs = []
            for column, value in enumerate(row, 1):
                if value:
                    chunks.append(_chunk(value, {'row': index, 'column': column, 'cell': f'R{index}C{column}'}))
                    heading = headers[column - 1] if column <= len(headers) else f'第{column}列'
                    pairs.append(heading + '=' + value)
            if pairs:
                chunks.append(_chunk(' | '.join(pairs), {'row': index}, kind='table_row'))
            _bounded(chunks)
    except csv.Error:
        service.fail(422, 'CSV 格式损坏，请重新保存后上传')
    return chunks, [], []


def read_docx(content):
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with _zip(content, 'word/document.xml') as archive:
        for name in archive.namelist():
            if name.endswith('.rels'):
                root = _xml(archive.read(name))
                if any(item.get('TargetMode') == 'External' for item in root):
                    service.fail(422, 'Word 含外部链接，请移除后上传')
        root = _xml(archive.read('word/document.xml'))
    def paragraph_text(paragraph):
        out = []
        for node in paragraph.iter():
            if node.tag == '{' + ns['w'] + '}t':
                out.append(node.text or '')
            elif node.tag in ('{' + ns['w'] + '}tab', '{' + ns['w'] + '}br'):
                out.append('\t' if node.tag.endswith('tab') else '\n')
        return ''.join(out)
    chunks = []
    for index, paragraph in enumerate(root.findall('.//w:p', ns), 1):
        text = paragraph_text(paragraph)
        if text.strip():
            chunks.append(_chunk(text, {'paragraph': index}))
        _bounded(chunks)
    for table_index, table in enumerate(root.findall('.//w:tbl', ns), 1):
        headers = []
        for row_index, row in enumerate(table.findall('w:tr', ns), 1):
            values = ['\n'.join(paragraph_text(p) for p in cell.findall('.//w:p', ns)) for cell in row.findall('w:tc', ns)]
            if row_index == 1:
                headers = values
            for column, value in enumerate(values, 1):
                if value:
                    chunks.append(_chunk(value, {'table': table_index, 'row': row_index, 'column': column}))
            if any(values):
                chunks.append(_chunk(' | '.join((headers[n] if n < len(headers) else f'第{n + 1}列') + '=' + value for n, value in enumerate(values)),
                                     {'table': table_index, 'row': row_index}, kind='table_row'))
            _bounded(chunks)
    return chunks, [], []


def read_pdf(content, source_id):
    try:
        import fitz
    except ImportError:
        service.fail(503, 'PDF 读取组件尚未安装，请联系本地维护人员；原资料未替换')
    if not content.startswith(b'%PDF-'):
        service.fail(422, '文件内容与 PDF 扩展名不符')
    chunks, pages, notes = [], [], []
    try:
        with fitz.open(stream=content, filetype='pdf') as document:
            if document.needs_pass:
                service.fail(422, 'PDF 已加密，请解除密码后上传')
            if not 1 <= document.page_count <= MAX_PAGES:
                service.fail(422, f'PDF 请拆分为每份 1 至 {MAX_PAGES} 页后上传')
            for index, page in enumerate(document, 1):
                text = page.get_text('text', sort=True)
                image_id = db.uid('image')
                scale = min(1.5, (4000000 / max(1, page.rect.width * page.rect.height)) ** .5)
                pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                stored = _write_once(f'pages/{source_id}_{index:04d}.png', pixmap.tobytes('png'))
                scanned = len(re.sub(r'\s', '', text)) < 24
                pages.append({'page': index, 'image_id': image_id, 'stored_name': stored, 'mime_type': 'image/png', 'scanned': scanned})
                if text.strip():
                    chunks.append(_chunk(text, {'page': index}, image_id=image_id))
                if scanned:
                    notes.append(f'第 {index} 页为扫描页或文字稀少，提取时必须提交该页原图，人工需核对页图')
                _bounded(chunks)
    except Exception as exc:
        from fastapi import HTTPException
        if isinstance(exc, HTTPException):
            raise
        service.fail(422, '无法读取 PDF，请重新保存有效文件后上传')
    return chunks, pages, notes


def read_image(content, source_id):
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('error', Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(content)) as image:
                if image.format not in ('PNG', 'JPEG', 'WEBP') or image.width * image.height > MAX_PIXELS or getattr(image, 'n_frames', 1) != 1:
                    service.fail(422, '请上传单张 PNG、JPG 或 WEBP 图片，像素不能超过 2400 万')
                image.load()
                # The normalized preview does not carry executable metadata;
                # the immutable original is stored separately by parse_file.
                preview = image.convert('RGB')
                preview.thumbnail((2400, 2400))
                buffer = io.BytesIO()
                preview.save(buffer, format='PNG')
                stored = _write_once(f'pages/{source_id}_0001.png', buffer.getvalue())
                mime = Image.MIME[image.format]
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombWarning, Image.DecompressionBombError):
        service.fail(422, '图片无法读取或像素过大，请重新导出后上传')
    return [], [{'page': 1, 'image_id': db.uid('image'), 'stored_name': stored, 'mime_type': 'image/png', 'scanned': True}], ['图片文字须由 AI 图像提取后逐项人工核对'], mime


def parse_file(filename, content, role):
    if not isinstance(filename, str) or not filename or len(filename) > 200 or '/' in filename or '\\' in filename:
        service.fail(422, '资料文件名格式不正确')
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_FILE_BYTES:
        service.fail(422, '资料文件应大于零且不超过 8 兆，请拆分后上传')
    extension = Path(filename).suffix.lower()
    kind = EXTENSIONS.get(extension)
    if not kind:
        service.fail(422, '仅支持 xlsx、csv、docx、pdf、png、jpg、jpeg、webp 资料')
    identity = db.uid('source')
    mime = MEDIA.get(kind)
    if kind == 'xlsx':
        chunks, pages, notes = read_xlsx(content)
    elif kind == 'csv':
        chunks, pages, notes = read_csv(content)
    elif kind == 'docx':
        chunks, pages, notes = read_docx(content)
    elif kind == 'pdf':
        chunks, pages, notes = read_pdf(content, identity)
    else:
        chunks, pages, notes, mime = read_image(content, identity)
    sha = hashlib.sha256(content).hexdigest()
    stored = _write_once(f'blobs/{sha}{extension}', content)
    return {'id': identity, 'name': filename, 'kind': kind, 'role': role, 'sha256': sha, 'size': len(content),
            'mime_type': mime, 'stored_name': stored, 'chunks': chunks, 'pages': pages, 'notes': notes,
            'created_at': db.now(), 'extracted': False}


def parse_text(name, text, role):
    if not isinstance(text, str) or not text.strip() or len(text) > 30000 or '\x00' in text:
        service.fail(422, '粘贴文字应为 1 至 30000 字，请分批添加')
    if not isinstance(name, str) or len(name) > 200:
        service.fail(422, '文字资料名称过长')
    content = text.encode('utf-8')
    sha = hashlib.sha256(content).hexdigest()
    stored = _write_once(f'blobs/{sha}.txt', content)
    chunks = [_chunk(paragraph, {'paragraph': index}) for index, paragraph in enumerate(text.splitlines(), 1) if paragraph.strip()]
    return {'id': db.uid('source'), 'name': name or '粘贴聊天原文', 'kind': 'text', 'role': role, 'sha256': sha,
            'size': len(content), 'mime_type': 'text/plain', 'stored_name': stored, 'chunks': _bounded(chunks), 'pages': [], 'notes': [],
            'created_at': db.now(), 'extracted': False}


def verify_original(source):
    path = storage_path(source['stored_name'])
    if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != source['sha256']:
        service.fail(409, '原资料文件缺失或内容已变动，不能放行，请重新核对资料')
    return path
