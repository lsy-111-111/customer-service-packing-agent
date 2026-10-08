"""Employee-only trial packing workflow, separate from retail stock and orders."""
import hashlib
import io
import json
import math
import re
import secrets
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from fastapi import APIRouter, Body, Depends, File, UploadFile
from fastapi.responses import Response
from . import db, service, packing_sku

SCHEMA = '''
CREATE TABLE IF NOT EXISTS packing_tasks(id TEXT PRIMARY KEY,number TEXT UNIQUE NOT NULL,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS packing_operations(actor_id TEXT NOT NULL,operation TEXT NOT NULL,request_id TEXT NOT NULL,fingerprint TEXT NOT NULL,data TEXT NOT NULL,PRIMARY KEY(actor_id,operation,request_id));
CREATE TABLE IF NOT EXISTS packing_labels(code TEXT PRIMARY KEY,task_id TEXT NOT NULL,box_id TEXT NOT NULL,revision INTEGER NOT NULL,active INTEGER NOT NULL DEFAULT 1);
CREATE INDEX IF NOT EXISTS packing_label_task ON packing_labels(task_id,active);
CREATE TABLE IF NOT EXISTS packing_documents(id TEXT PRIMARY KEY,task_id TEXT NOT NULL,data TEXT NOT NULL,pdf BLOB NOT NULL);
CREATE TABLE IF NOT EXISTS packing_packages(product_id TEXT PRIMARY KEY,version INTEGER NOT NULL,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS packing_package_previews(id TEXT PRIMARY KEY,actor_id TEXT NOT NULL,product_id TEXT NOT NULL,base_fingerprint TEXT NOT NULL,data TEXT NOT NULL,expires_at TEXT NOT NULL,consumed INTEGER NOT NULL DEFAULT 0);
'''
STAGES = {'draft': '待记录实装', 'approved': '审核已通过', 'dispatched': '已模拟出库', 'archived': '已归档', 'cancelled': '已取消'}
HEADERS = ['商品名称', '商品规格', '发货数量', '每箱数量', '批次', '生产日期', '到期日期', '商品条码', '套装及配件说明', '序列号']


def init_schema(c):
    for statement in SCHEMA.split(';'):
        if statement.strip():
            c.execute(statement)
    packing_sku.init_schema(c)


def _staff(user):
    if user.get('role') != 'staff':
        service.fail(403, '请由员工进入装箱工作台')


def _text(value, title, limit=200):
    if value is None:
        return ''
    if not isinstance(value, str) or len(value.strip()) > limit or any(ord(ch) < 32 and ch not in '\n\t' for ch in value):
        service.fail(422, f'{title}内容过长或格式不正确')
    return value.strip()


def _number(value, title, maximum=100000, minimum=1):
    if isinstance(value, bool) or not isinstance(value, int) or not minimum <= value <= maximum:
        service.fail(422, f'{title}请填写 {minimum} 至 {maximum} 的整数')
    return value


def _date(value, title):
    value = _text(value, title, 10)
    if value:
        try:
            if date.fromisoformat(value).isoformat() != value:
                raise ValueError()
        except ValueError:
            service.fail(422, f'{title}请填写完整年月日')
    return value


def _serials(value):
    if isinstance(value, str):
        value = re.split(r'[,，、;；\s]+', value.strip()) if value.strip() else []
    if not isinstance(value, list) or len(value) > 100000:
        service.fail(422, '请逐个填写商品序列号，用逗号或换行分隔')
    values = [_text(item, '序列号', 80) for item in value]
    if any(not item for item in values) or len(set(values)) != len(values):
        service.fail(422, '序列号不能为空或重复，请核对商品')
    return values


REQUIREMENTS = ('require_batch', 'require_expiry', 'require_serials')


def _fingerprint(data):
    return hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def _package(c, product_id, product=None):
    if product is None:
        product = json.loads(service.one(c, 'SELECT data FROM products WHERE id=?', (product_id,), '商品不存在，请重新选择')[0])
    row = c.execute('SELECT version,data FROM packing_packages WHERE product_id=?', (product_id,)).fetchone()
    saved = json.loads(row['data']) if row else {}
    parts = saved.get('components', []) if row else [
        {'name': item['name'], 'quantity': item['quantity']}
        for item in product.get('package_items', []) if item.get('kind') == 'part']
    result = {'product_id': product_id, 'product_name': product['name'], 'version': row['version'] if row else 0,
              'source': 'packing' if row else 'product',
              'source_label': '已保存的装箱包装资料' if row else '沿用店内商品包装清单',
              'components': parts, 'requirements': {key: saved.get('requirements', {}).get(key, False) for key in REQUIREMENTS},
              'sku_settings': packing_sku.for_product(c, product_id)}
    result['fingerprint'] = _fingerprint(result)
    return result


def list_packages(user):
    _staff(user)
    with db.connect() as c:
        return [_package(c, row['id'], json.loads(row['data'])) for row in c.execute('SELECT id,data FROM products ORDER BY rowid')]


def catalog(user):
    _staff(user)
    with db.connect() as c:
        products = [service._product(c, row['id']) for row in c.execute('SELECT id FROM products ORDER BY rowid')]
        for product in products:
            for sku in product['skus']:
                state = packing_sku.current(c, sku['id'])
                sku.update(packing_enabled=state['packing_enabled'], product_barcode=state['product_barcode'], packing_version=state['version'])
        return products


def preview_package(user, data):
    _staff(user)
    with db.transaction() as c:
        current = _package(c, _text(data.get('product_id'), '商品', 100))
        if type(data.get('expected_version')) is not int or data['expected_version'] != current['version']:
            service.fail(409, '包装资料已更新，请重新选择商品后核对')
        parts = data.get('components')
        if not isinstance(parts, list) or len(parts) > 30:
            service.fail(422, '包装资料最多填写 30 种配件；无需配件时可以留空')
        clean, names = [], set()
        for index, part in enumerate(parts, 1):
            if not isinstance(part, dict):
                service.fail(422, f'第 {index} 项配件资料不完整')
            name = _text(part.get('name'), f'第 {index} 项配件名称', 80)
            if not name or name.casefold() in names:
                service.fail(422, '配件名称不能为空或重复，请逐项核对')
            names.add(name.casefold())
            clean.append({'name': name, 'quantity': _number(part.get('quantity'), f'第 {index} 项配件每套用量', 10000)})
        requirements = data.get('requirements', {})
        if not isinstance(requirements, dict) or any(type(requirements.get(key, False)) is not bool for key in REQUIREMENTS):
            service.fail(422, '请勾选适用的批次、效期和序列号要求')
        candidate = {'product_id': current['product_id'], 'product_name': current['product_name'], 'components': clean,
                     'requirements': {key: requirements.get(key, False) for key in REQUIREMENTS}}
        reason = _text(data.get('reason'), 'SKU 资料核对依据与原因', 500)
        candidate.update(sku_settings=packing_sku.proposed(current['sku_settings'], data.get('sku_settings'), reason), reason=reason)
        identity = db.uid('packpreview')
        expires = (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat()
        c.execute('INSERT INTO packing_package_previews VALUES (?,?,?,?,?,?,0)',
                  (identity, user['id'], current['product_id'], current['fingerprint'], service.js(candidate), expires))
        return {'preview_id': identity, 'before': current, 'candidate': candidate, 'expires_at': expires,
                'message': '请核对配件、默认追溯要求及每个 SKU 的条码和装箱启停。旧任务保留资料快照；未完成任务在审核和放行前重查当前 SKU 状态。'}


def confirm_package(user, data):
    _staff(user)
    with db.transaction() as c:
        def action():
            row = service.one(c, 'SELECT * FROM packing_package_previews WHERE id=? AND actor_id=?',
                              (_text(data.get('preview_id'), '核对记录', 100), user['id']), '请先预览本次包装资料')
            if row['consumed'] or datetime.fromisoformat(row['expires_at']) < datetime.now(timezone.utc):
                service.fail(409, '这次核对记录已使用或过期，请重新预览包装资料')
            current = _package(c, row['product_id'])
            if row['base_fingerprint'] != current['fingerprint']:
                service.fail(409, '包装资料已被更新，请重新选择商品并预览后确认')
            candidate = json.loads(row['data'])
            candidate.update(updated_at=db.now(), operator=user['name'])
            packing_sku.save(c, user, candidate.get('sku_settings', current['sku_settings']), candidate.get('reason', ''))
            c.execute('INSERT INTO packing_packages VALUES (?,?,?) ON CONFLICT(product_id) DO UPDATE SET version=excluded.version,data=excluded.data',
                      (row['product_id'], current['version'] + 1, service.js(candidate)))
            c.execute('UPDATE packing_package_previews SET consumed=1 WHERE id=?', (row['id'],))
            service.audit(c, user, 'packing_package_save', row['product_id'], current, candidate)
            return _package(c, row['product_id'])
        return _operation(c, user, 'packing_package_save', data, action)


def _normalize(c, data, check_packaging=True):
    lines = data.get('lines')
    if not isinstance(lines, list) or not 1 <= len(lines) <= 50:
        service.fail(422, '请添加 1 至 50 行发货商品')
    result = {'destination': _text(data.get('destination'), '收货仓库'),
              'supplier': _text(data.get('supplier'), '发货单位'),
              'platform_order': _text(data.get('platform_order'), '平台单号', 80),
              'ship_date': _date(data.get('ship_date'), '计划发货日期'),
              'note': _text(data.get('note'), '发货备注', 500), 'lines': [], 'boxes': [], 'materials': []}
    serials_seen = set()
    for n, item in enumerate(lines, 1):
        if not isinstance(item, dict):
            service.fail(422, f'第 {n} 行商品资料不完整')
        sid = _text(item.get('sku_id'), '商品编号', 100)
        row = service.one(c, 'SELECT skus.*,products.data product FROM skus JOIN products ON products.id=skus.product_id WHERE skus.id=?', (sid,), f'第 {n} 行请选择店内已有商品规格')
        product = json.loads(row['product'])
        package = _package(c, row['product_id'], product)
        if check_packaging and item.get('packaging_fingerprint') and item['packaging_fingerprint'] != package['fingerprint']:
            service.fail(409, f'第 {n} 行包装资料已更新，请重新计算并核对分箱')
        qty = _number(item.get('quantity'), f'第 {n} 行发货数量')
        per = _number(item.get('per_carton'), f'第 {n} 行每箱数量', 10000)
        produced = _date(item.get('produced_at'), '生产日期')
        expires = _date(item.get('expires_at'), '到期日期')
        if produced and expires and produced > expires:
            service.fail(422, f'第 {n} 行到期日期不能早于生产日期')
        serials = _serials(item.get('serial_numbers', []))
        if serials and len(serials) != qty:
            service.fail(422, f'第 {n} 行序列号数量须与发货数量一致')
        if serials_seen.intersection(serials):
            service.fail(422, '不同商品行存在重复序列号，请核对')
        serials_seen.update(serials)
        sku_state = packing_sku.current(c, sid)
        supplied_barcode = item.get('product_barcode')
        barcode = packing_sku.barcode(sku_state['product_barcode'] if supplied_barcode in (None, '') else supplied_barcode)
        safety_review = packing_sku.review(sku_state, barcode, item.get('sku_confirmation'))
        line = {'sku_id': sid, 'product_name': product['name'], 'specification': row['label'],
                'external_code': _text(item.get('external_code'), '客户商品编码', 200),
                'quantity': qty, 'per_carton': per, 'batch': _text(item.get('batch'), '批次', 80),
                'produced_at': produced, 'expires_at': expires, 'product_barcode': barcode,
                'kit_note': _text(item.get('kit_note'), '套装及配件说明', 500),
                'package_items': ([{'kind': 'main', 'name': product['name'], 'quantity': 1}] +
                                  [{'kind': 'part', **part} for part in package['components']]) if package['source'] == 'packing' else product.get('package_items', []),
                'serial_numbers': serials, 'packaging_snapshot': package, 'packaging_fingerprint': package['fingerprint']}
        if safety_review:
            line['sku_safety_review'] = safety_review
        for key in REQUIREMENTS:
            if key in item and type(item[key]) is not bool:
                service.fail(422, f'第 {n} 行请勾选适用的追溯要求')
            line[key] = item.get(key, package['requirements'][key])
        result['lines'].append(line)
        for component in line['package_items']:
            if component.get('kind') == 'part':
                if component['quantity'] * qty > 1000000:
                    service.fail(422, '配件需备数量过大，请拆成多个装箱任务')
                result['materials'].append({'id': db.uid('material'), 'line_index': n - 1,
                                            'name': component['name'], 'per_item': component['quantity'],
                                            'planned_quantity': component['quantity'] * qty, 'actual_quantity': None})
        for offset in range(0, qty, per):
            if len(result['boxes']) >= 300:
                service.fail(422, '一次装箱任务最多 300 箱，请分成多个任务')
            result['boxes'].append({'id': db.uid('box'), 'number': len(result['boxes']) + 1,
                                    'line_index': n - 1, 'planned_quantity': min(per, qty - offset),
                                    'serial_numbers': serials[offset:offset + per], 'actual_quantity': None,
                                    'net_weight': None, 'gross_weight': None, 'actual_at': None,
                                    'code': None, 'labeled_at': None, 'scanned_at': None})
    result['total_quantity'] = sum(item['quantity'] for item in result['lines'])
    result['total_boxes'] = len(result['boxes'])
    result['warnings'] = ['试用样张，仓库要求待核对。当前未接入京东供应商系统。', '每箱只装同一商品规格及批次；套装内配件不另计成品箱。']
    if not result['destination']:
        result['warnings'].append('尚未填写收货仓库，正式审核前需要补齐。')
    return result


def preview(user, data):
    _staff(user)
    with db.connect() as c:
        return _normalize(c, data, check_packaging=False)


def _get(c, identity):
    return json.loads(service.one(c, 'SELECT data FROM packing_tasks WHERE id=?', (identity,), '装箱任务不存在')[0])


def _save(c, task):
    task['stage_label'] = STAGES[task['stage']]
    c.execute('UPDATE packing_tasks SET data=? WHERE id=?', (service.js(task), task['id']))


def _history(task, user, action, note=''):
    task['updated_at'] = db.now()
    task['history'].append({'at': task['updated_at'], 'operator': user['name'], 'action': action, 'note': note})


def _operation(c, user, operation, data, callback):
    key = _text(data.get('request_id'), '提交凭证', 200)
    if not key:
        service.fail(422, '页面提交凭证已丢失，请刷新后重新操作')
    fingerprint = hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    row = c.execute('SELECT fingerprint,data FROM packing_operations WHERE actor_id=? AND operation=? AND request_id=?', (user['id'], operation, key)).fetchone()
    if row:
        if row[0] != fingerprint:
            service.fail(409, '这次提交的内容已变化，请刷新后重新核对')
        return json.loads(row[1])
    result = callback()
    c.execute('INSERT INTO packing_operations VALUES (?,?,?,?,?)', (user['id'], operation, key, fingerprint, service.js(result)))
    return result


def create_task(user, data):
    _staff(user)
    with db.transaction() as c:
        def action():
            task = _normalize(c, data)
            task.update(id=db.uid('packing'), number='ZX' + date.today().strftime('%Y%m%d') + secrets.token_hex(3).upper(),
                        version=1, revision=1, stage='draft', stage_label=STAGES['draft'], approval=None,
                        created_at=db.now(), creator=user['name'], history=[])
            from . import packing_sources, packing_workflow
            task['layout'] = packing_workflow.layout(data.get('layout', {}))
            task['demo'] = data.get('demo') is True
            packing_sources.bind_task(c, user, task, data)
            _history(task, user, '核对后建立装箱任务')
            c.execute('INSERT INTO packing_tasks VALUES (?,?,?)', (task['id'], task['number'], service.js(task)))
            service.audit(c, user, 'packing_create', task['id'], None, {'number': task['number']})
            return task
        return _operation(c, user, 'packing_create', data, action)


def list_tasks(user):
    _staff(user)
    with db.connect() as c:
        tasks = [json.loads(row[0]) for row in c.execute('SELECT data FROM packing_tasks ORDER BY rowid DESC LIMIT 300')]
        return [{key: task[key] for key in ('id', 'number', 'stage', 'stage_label', 'total_quantity', 'total_boxes', 'destination', 'created_at', 'updated_at', 'creator')} for task in tasks]


def get_task(user, identity):
    _staff(user)
    with db.connect() as c:
        task = _get(c, identity)
        task['documents'] = [json.loads(row[0]) for row in c.execute('SELECT data FROM packing_documents WHERE task_id=? ORDER BY rowid DESC', (identity,))]
        task['sku_checks'] = packing_sku.task_checks(c, task)
        return task


def _version(task, data):
    if type(data.get('version')) is not int or data.get('version') != task['version']:
        service.fail(409, '任务已被更新，请重新打开并核对最新内容')


def _editable(task):
    if task['stage'] in ('dispatched', 'archived', 'cancelled'):
        service.fail(409, '已模拟出库或归档的任务不能再改动，请新建任务')


def _invalidate(c, task):
    c.execute('UPDATE packing_labels SET active=0 WHERE task_id=?', (task['id'],))
    for box in task['boxes']:
        box.update(code=None, labeled_at=None, scanned_at=None)
    task.update(approval=None, stage='draft', revision=task['revision'] + 1)


def task_action(user, identity, data):
    _staff(user)
    with db.transaction() as c:
        def action():
            task = _get(c, identity)
            _version(task, data)
            kind = data.get('action')
            from . import packing_sources, packing_workflow
            if kind == 'revise':
                _editable(task)
                reason = _text(data.get('reason'), '修改原因', 300)
                if not reason:
                    service.fail(422, '请填写修改原因，便于追溯旧标签')
                replacement = _normalize(c, data)
                packing_workflow.save_revision(c, task, reason)
                _invalidate(c, task)
                task.update(replacement)
                if 'layout' in data:
                    task['layout'] = packing_workflow.layout(data['layout'])
                packing_sources.bind_task(c, user, task, data)
                _history(task, user, '修改装箱资料，原审核与标签作废', reason)
            elif kind == 'actual':
                _editable(task)
                if task['approval']:
                    service.fail(409, '审核已锁定，请先修改装箱资料并说明原因')
                rows = data.get('boxes')
                if not isinstance(rows, list) or len(rows) != len(task['boxes']):
                    service.fail(422, '请逐箱核对并提交全部实装记录')
                found = {row.get('id'): row for row in rows if isinstance(row, dict)}
                if len(found) != len(task['boxes']):
                    service.fail(422, '实装箱号存在重复或遗漏，请刷新核对')
                for box in task['boxes']:
                    item = found.get(box['id'])
                    if not item:
                        service.fail(422, '实装记录与当前箱号不符，请刷新核对')
                    actual = _number(item.get('actual_quantity'), f'第 {box["number"]} 箱实装数量', 10000, 0)
                    weights = []
                    for field, title in (('net_weight', '净重'), ('gross_weight', '毛重')):
                        value = item.get(field)
                        if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 < value <= 10000):
                            service.fail(422, f'第 {box["number"]} 箱{title}请填写大于零的公斤数，未知可留空')
                        weights.append(value)
                    if all(value is not None for value in weights) and weights[1] < weights[0]:
                        service.fail(422, f'第 {box["number"]} 箱毛重不能小于净重')
                    box.update(actual_quantity=actual, net_weight=weights[0], gross_weight=weights[1], actual_at=db.now(), actual_operator=user['name'])
                materials = data.get('materials', [])
                if not isinstance(materials, list):
                    service.fail(422, '请核对配件备料数量')
                material_map = {item.get('id'): item for item in materials if isinstance(item, dict)}
                if len(material_map) != len(materials) or set(material_map) != {item['id'] for item in task.get('materials', [])}:
                    service.fail(422, '配件备料记录存在重复或遗漏，请核对全部配件')
                for material in task.get('materials', []):
                    material['actual_quantity'] = _number(material_map[material['id']].get('actual_quantity'), material['name'] + '实备数量', 1000000, 0)
                _history(task, user, '已逐箱记录实装数量与重量')
            elif kind == 'approve':
                _editable(task)
                if task['approval']:
                    service.fail(409, '当前任务已经审核通过')
                if not task['destination'] or not task['supplier'] or not task['ship_date']:
                    service.fail(422, '审核前请补齐发货单位、收货仓库和计划发货日期')
                if any(box['actual_quantity'] != box['planned_quantity'] or not box['actual_at'] for box in task['boxes']):
                    service.fail(422, '请先逐箱记录实装数量，并处理实装与计划的差异')
                for line in task['lines']:
                    if line.get('require_batch') and not line['batch']:
                        service.fail(422, f'{line["product_name"]}已要求批次追溯，请补齐批次')
                    if line.get('require_expiry') and (not line['produced_at'] or not line['expires_at']):
                        service.fail(422, f'{line["product_name"]}已要求效期管理，请补齐生产和到期日期')
                    if line.get('require_serials') and len(line['serial_numbers']) != line['quantity']:
                        service.fail(422, f'{line["product_name"]}已要求逐件追溯，请补齐全部序列号')
                    if line['expires_at'] and line['expires_at'] <= task['ship_date']:
                        service.fail(422, '存在计划发货日已到期的商品，请核对效期')
                    if line['produced_at'] and line['produced_at'] > task['ship_date']:
                        service.fail(422, '生产日期不能晚于计划发货日期')
                if any(item['actual_quantity'] != item['planned_quantity'] for item in task.get('materials', [])):
                    service.fail(422, '请核对套装配件备料，实备数量须与需求一致')
                if data.get('confirmed') is not True:
                    service.fail(422, '请确认已核对实装、商品批次及收货资料')
                packing_workflow.check_approval(c, task)
                preview_document_id = packing_workflow.approval_preview(c, task, data)
                task.update(stage='approved', approval={'operator': user['name'], 'at': db.now(), 'revision': task['revision'],
                                                       'preview_document_id': preview_document_id,
                                                       'content_fingerprint': packing_workflow.content_fingerprint(task)})
                _history(task, user, '审核通过，装箱资料已锁定', '仅本地审核；京东仓库要求仍待核对')
            elif kind == 'label':
                _editable(task)
                packing_workflow.check_approval(c, task)
                boxes = _selected(task, data)
                if not task['approval'] or any(not box['code'] for box in boxes):
                    service.fail(409, '请先审核并生成当前箱唛，再确认贴标')
                for box in boxes:
                    box['labeled_at'] = db.now()
                _history(task, user, '已确认贴好当前标签', '箱号：' + '、'.join(str(box['number']) for box in boxes))
            elif kind == 'dispatch':
                packing_workflow.check_approval(c, task)
                if task['stage'] != 'approved' or not all(box['code'] and box['labeled_at'] and box['scanned_at'] for box in task['boxes']):
                    service.fail(409, '请完成审核、贴标和全部箱子的扫码核验后再模拟出库')
                task.update(stage='dispatched', dispatched_at=db.now())
                _history(task, user, '已模拟出库', '本地演练，不扣库存、不创建真实物流单')
            elif kind == 'archive':
                if task['stage'] != 'dispatched':
                    service.fail(409, '请先完成模拟出库再归档')
                task.update(stage='archived', archived_at=db.now())
                _history(task, user, '任务已归档')
            else:
                service.fail(422, '请选择有效的装箱操作')
            task['version'] += 1
            _save(c, task)
            service.audit(c, user, 'packing_' + kind, identity, None, {'number': task['number'], 'stage': task['stage']})
            return task
        return _operation(c, user, 'packing_action:' + identity, data, action)


def _selected(task, data):
    chosen = data.get('box_ids')
    if chosen is None:
        return task['boxes']
    if not isinstance(chosen, list) or not chosen or any(not isinstance(value, str) for value in chosen) or len(set(chosen)) != len(chosen):
        service.fail(422, '请选择需要操作的箱子')
    boxes = [box for box in task['boxes'] if box['id'] in chosen]
    if len(boxes) != len(chosen):
        service.fail(422, '所选箱子已变化，请刷新后核对')
    return boxes


def scan(user, data):
    _staff(user)
    with db.transaction() as c:
        def action():
            code = _text(data.get('code'), '箱子核验号', 80)
            row = service.one(c, 'SELECT * FROM packing_labels WHERE code=?', (code,), '没有找到这个箱子核验号，请扫描内部核验标签')
            task = _get(c, row['task_id'])
            if data.get('task_id') and data['task_id'] != task['id']:
                service.fail(409, '这张标签属于另一项装箱任务，请核对后打开对应任务')
            if not row['active'] or row['revision'] != task['revision'] or not task['approval']:
                service.fail(409, '这张标签已作废，请换用最新补印的标签')
            if task['stage'] != 'approved':
                service.fail(409, '任务已结束，不能继续扫码出库')
            from .packing_workflow import check_approval
            check_approval(c, task)
            box = next((item for item in task['boxes'] if item['id'] == row['box_id']), None)
            if not box or box['code'] != code or not box['labeled_at']:
                service.fail(409, '请先确认该箱已贴好当前标签')
            duplicate = bool(box['scanned_at'])
            if not duplicate:
                box['scanned_at'] = db.now()
                task['version'] += 1
                _history(task, user, '扫码核验通过', f'第 {box["number"]} 箱')
                _save(c, task)
            return {'task': task, 'duplicate': duplicate, 'message': f'第 {box["number"]} 箱' + ('此前已核验，本次不重复计数' if duplicate else '核验通过')}
        return _operation(c, user, 'packing_scan', data, action)


def create_document(user, identity, data):
    _staff(user)
    with db.transaction() as c:
        def action():
            task = _get(c, identity)
            _version(task, data)
            _editable(task)
            mode = data.get('mode', 'draft')
            if mode not in ('draft', 'formal'):
                service.fail(422, '请选择草稿或审核后箱唛')
            if mode == 'formal' and not task['approval']:
                service.fail(409, '请先记录实装并完成审核，当前只能生成草稿')
            boxes = _selected(task, data)
            reprint = data.get('reprint') is True
            reason = _text(data.get('reason'), '补印原因', 300)
            if reprint and (mode != 'formal' or not reason or any(not box['code'] for box in boxes)):
                service.fail(422, '补印请选已有标签的箱子，并填写补印原因')
            if mode == 'formal':
                for box in boxes:
                    if reprint and box['code']:
                        c.execute('UPDATE packing_labels SET active=0 WHERE code=?', (box['code'],))
                        box.update(code=None, labeled_at=None, scanned_at=None)
                    if not box['code']:
                        box['code'] = 'BX' + secrets.token_hex(8).upper()
                        c.execute('INSERT INTO packing_labels VALUES (?,?,?,?,1)', (box['code'], identity, box['id'], task['revision']))
            paper = data.get('paper', 'a4')
            kind = data.get('kind', 'all')
            if paper not in ('a4', 'label') or kind not in ('main', 'side', 'internal', 'all'):
                service.fail(422, '请选择支持的纸张和标签内容')
            width = _number(data.get('width_mm', 100), '标签宽度', 200)
            height = _number(data.get('height_mm', 150), '标签高度', 270)
            if width < 100 or height < 150:
                service.fail(422, '试用模板至少需要宽 100 毫米、高 150 毫米，避免内容过小')
            from . import packing_workflow, packing_sources
            selected_layout = packing_workflow.layout({'paper': paper, 'kind': kind, 'width_mm': width, 'height_mm': height,
                                                       'show_note': data.get('show_note', task.get('layout', {}).get('show_note', True))})
            if mode == 'formal':
                packing_workflow.check_approval(c, task)
                approved_layout = packing_workflow.layout(task.get('layout', {}))
                if any(selected_layout[k] != approved_layout[k] for k in ('paper', 'width_mm', 'height_mm', 'show_note')):
                    service.fail(409, '打印版式与已审核内容不一致，请保存版式、重新预览并审核')
                if hasattr(packing_sources, 'validate_document_rules'):
                    packing_sources.validate_document_rules(c, task, selected_layout)
            from .packing_pdf import render
            pages = []
            pdf = render(task, boxes, mode, paper, kind, width, height, pages=pages, show_note=selected_layout['show_note'])
            doc = {'id': db.uid('packingdoc'), 'task_id': identity, 'number': task['number'], 'created_at': db.now(),
                   'operator': user['name'], 'mode': mode, 'label': '草稿箱唛' if mode == 'draft' else '审核后试用箱唛',
                   'revision': task['revision'], 'paper': paper, 'kind': kind, 'box_numbers': [box['number'] for box in boxes],
                   'codes': [box['code'] for box in boxes] if mode == 'formal' else [], 'reprint': reprint,
                   'width_mm': width, 'height_mm': height, 'layout': selected_layout, 'pages': pages,
                   'content_fingerprint': packing_workflow.content_fingerprint(task)}
            c.execute('INSERT INTO packing_documents VALUES (?,?,?,?)', (doc['id'], identity, service.js(doc), pdf))
            _history(task, user, '补印标签，所选旧标签作废' if reprint else '生成' + doc['label'], reason)
            task['version'] += 1
            _save(c, task)
            return {'document': doc, 'task': task}
        return _operation(c, user, 'packing_document:' + identity, data, action)


def download_document(user, identity):
    _staff(user)
    with db.connect() as c:
        row = service.one(c, 'SELECT data,pdf FROM packing_documents WHERE id=?', (identity,), '箱唛文件不存在')
        meta = json.loads(row['data'])
        if meta['mode'] == 'formal':
            task = _get(c, meta['task_id'])
            from .packing_workflow import check_approval
            check_approval(c, task)
            current_codes = {box['code'] for box in task['boxes']}
            if not task['approval'] or meta['revision'] != task['revision'] or any(code not in current_codes for code in meta['codes']):
                service.fail(409, '该版标签已失效，请下载当前箱唛')
        return bytes(row['pdf']), meta


def assist(user, data):
    _staff(user)
    text = _text(data.get('text'), '装箱描述', 2000)
    per = re.search(r'每箱\s*(?:装|放)?\s*(\d+)\s*(?:件|套|个|盒|台)?', text)
    without_per = text[:per.start()] + text[per.end():] if per else text
    total = re.search(r'(?:总共|总计|合计|发货|一共|共|总量)\s*(\d+)\s*(?:件|套|个|盒|台)?', without_per) or re.search(r'(\d+)\s*(?:件|套|个|盒|台)', without_per)
    matches = []
    for product in service.products():
        names = [product['name'], product.get('model', '')] + product.get('aliases', [])
        if any(name and name.lower() in text.lower() for name in names):
            for sku in product['skus']:
                matches.append({'sku_id': sku['id'], 'name': product['name'], 'specification': sku['label']})
    exact = [item for item in matches if item['specification'] in text or item['sku_id'] in text]
    matches = exact or matches
    questions = []
    if len(matches) != 1:
        questions.append('请在下方选择准确的商品规格。')
    if not total:
        questions.append('请补充总发货数量。')
    if not per:
        questions.append('请补充每箱装多少件。')
    return {'quantity': int(total.group(1)) if total else None, 'per_carton': int(per.group(1)) if per else None,
            'sku_id': matches[0]['sku_id'] if len(matches) == 1 else None, 'matches': matches,
            'questions': questions, 'message': '已整理描述，请在表单核对；不会直接建立任务。'}


def import_excel(user, content):
    _staff(user)
    if len(content) > 2 * 1024 * 1024:
        service.fail(422, '表格不能超过 2 兆，请分批导入')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                service.fail(422, '表格解压后过大，请使用标准模板重新填写')
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, ValueError, KeyError, OSError):
        service.fail(422, '无法读取表格，请上传标准工作簿格式的装箱模板')
    try:
        if '装箱明细' not in workbook.sheetnames:
            service.fail(422, '没有找到“装箱明细”工作表，请使用标准模板')
        sheet = workbook['装箱明细']
        if (sheet.max_row is not None and sheet.max_row > 51) or (sheet.max_column is not None and sheet.max_column > 10):
            service.fail(422, '装箱明细最多 50 行、10 列，请按标准模板分批填写')
        # Some valid workbooks omit dimensions. Iteration discovers real rows;
        # stop on the first excess row instead of silently truncating a sparse sheet.
        rows = sheet.iter_rows()
        first = next(rows, ())
        headers = [str(cell.value or '').strip() for cell in first[:10]]
        if headers != HEADERS or any(cell.value is not None for cell in first[10:]):
            service.fail(422, '表头与标准模板不一致，请使用下载的模板并保持列名和顺序')
        lines = []
        for rn, cells in enumerate(rows, 2):
            if rn > 51:
                service.fail(422, '一次最多导入 50 行商品，请分批填写')
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == 'f' for cell in cells):
                service.fail(422, f'第 {rn} 行包含公式，请复制为实际数值后再导入')
            if any(cell.value is not None for cell in cells[10:]):
                service.fail(422, f'第 {rn} 行多出一列内容，请按模板填写')
            values = [cell.value for cell in cells[:10]]
            values.extend([None] * (10 - len(values)))
            for index, title in ((7, '商品条码'), (9, '序列号')):
                if values[index] is not None and not isinstance(values[index], str):
                    service.fail(422, f'第 {rn} 行{title}请按文本填写，避免丢失前导零或长编号精度')
            for index in (2, 3):
                if isinstance(values[index], float) and values[index].is_integer():
                    values[index] = int(values[index])
            for index in (5, 6):
                if hasattr(values[index], 'date'):
                    values[index] = values[index].date().isoformat()
                elif isinstance(values[index], date):
                    values[index] = values[index].isoformat()
            def string(index):
                return '' if values[index] is None else str(values[index]).strip()
            with db.connect() as c:
                matches = []
                for sku in c.execute('SELECT skus.id,skus.label,products.data FROM skus JOIN products ON products.id=skus.product_id'):
                    if json.loads(sku['data'])['name'] == string(0) and sku['label'] == string(1):
                        matches.append(sku['id'])
            if len(matches) != 1:
                service.fail(422, f'第 {rn} 行商品名称或规格未能唯一匹配，请按店内名称填写')
            lines.append({'sku_id': matches[0], 'quantity': values[2], 'per_carton': values[3],
                          'batch': string(4), 'produced_at': string(5), 'expires_at': string(6),
                          'product_barcode': string(7), 'kit_note': string(8), 'serial_numbers': string(9)})
        with db.connect() as c:
            normalized = _normalize(c, {'lines': lines})
        return {'lines': normalized['lines'], 'message': f'已读取 {len(lines)} 行，请补齐发货资料后统一核对。'}
    finally:
        workbook.close()


def build_router(user_dependency):
    router = APIRouter(prefix='/api/packing')

    @router.get('/catalog')
    def catalog_route(u=Depends(user_dependency)):
        return catalog(u)

    @router.get('/packages')
    def packages_route(u=Depends(user_dependency)):
        return list_packages(u)

    @router.post('/packages/preview')
    def package_preview_route(data: dict = Body(...), u=Depends(user_dependency)):
        return preview_package(u, data)

    @router.post('/packages/confirm')
    def package_confirm_route(data: dict = Body(...), u=Depends(user_dependency)):
        return confirm_package(u, data)

    @router.get('/tasks')
    def listing(u=Depends(user_dependency)):
        return list_tasks(u)

    @router.post('/preview')
    def preview_route(data: dict = Body(...), u=Depends(user_dependency)):
        return preview(u, data)

    @router.post('/assist')
    def assist_route(data: dict = Body(...), u=Depends(user_dependency)):
        return assist(u, data)

    @router.post('/import')
    async def import_route(file: UploadFile = File(...), u=Depends(user_dependency)):
        _staff(u)
        if not file.filename or not file.filename.lower().endswith('.xlsx'):
            service.fail(422, '请上传标准工作簿文件，文件名以 .xlsx 结尾')
        content = await file.read(2 * 1024 * 1024 + 1)
        return import_excel(u, content)

    @router.post('/tasks')
    def create_route(data: dict = Body(...), u=Depends(user_dependency)):
        return create_task(u, data)

    @router.get('/tasks/{identity}')
    def detail_route(identity: str, u=Depends(user_dependency)):
        return get_task(u, identity)

    @router.post('/tasks/{identity}/actions')
    def action_route(identity: str, data: dict = Body(...), u=Depends(user_dependency)):
        return task_action(u, identity, data)

    @router.post('/tasks/{identity}/documents')
    def document_route(identity: str, data: dict = Body(...), u=Depends(user_dependency)):
        return create_document(u, identity, data)

    @router.get('/documents/{identity}')
    def download_route(identity: str, u=Depends(user_dependency)):
        content, meta = download_document(u, identity)
        return Response(content, media_type='application/pdf', headers={'Content-Disposition': f'inline; filename="{meta["number"]}.pdf"', 'Cache-Control': 'private, no-store'})

    @router.post('/scan')
    def scan_route(data: dict = Body(...), u=Depends(user_dependency)):
        return scan(u, data)

    return router
