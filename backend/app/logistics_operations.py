"""Local, sourced shipment evidence and internal follow-up. No carrier or fulfillment calls."""
import csv
import hashlib
import io
import json
import re
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from . import db, service, work_operations as work

SCHEMA = '''
CREATE TABLE IF NOT EXISTS logistics_shipments(id TEXT PRIMARY KEY,natural_key TEXT NOT NULL UNIQUE,order_id TEXT NOT NULL,user_id TEXT NOT NULL,source_kind TEXT NOT NULL,version INTEGER NOT NULL,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS logistics_events(id TEXT PRIMARY KEY,shipment_id TEXT NOT NULL,dedupe_key TEXT NOT NULL UNIQUE,source_event_key TEXT UNIQUE,effective INTEGER NOT NULL,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS logistics_imports(id TEXT PRIMARY KEY,actor_id TEXT NOT NULL,fingerprint TEXT NOT NULL,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS logistics_previews(id TEXT PRIMARY KEY,actor_id TEXT NOT NULL,expires_at TEXT NOT NULL,consumed INTEGER NOT NULL,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS logistics_rules(id INTEGER PRIMARY KEY CHECK(id=1),version INTEGER NOT NULL,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS logistics_rule_history(id TEXT PRIMARY KEY,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS logistics_exceptions(id TEXT PRIMARY KEY,shipment_id TEXT NOT NULL,kind TEXT NOT NULL,status TEXT NOT NULL,version INTEGER NOT NULL,data TEXT NOT NULL);
CREATE UNIQUE INDEX IF NOT EXISTS logistics_exception_open ON logistics_exceptions(shipment_id,kind) WHERE status!='closed';
CREATE INDEX IF NOT EXISTS logistics_event_shipment ON logistics_events(shipment_id);
CREATE TABLE IF NOT EXISTS logistics_judgements(id TEXT PRIMARY KEY,shipment_id TEXT NOT NULL,exception_id TEXT,data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS logistics_exception_history(id TEXT PRIMARY KEY,exception_id TEXT NOT NULL,data TEXT NOT NULL);
'''
EVENT_TYPES = ['已交寄待揽收', '已揽收', '运输中', '派送中', '已签收', '拒收', '破损', '退回', '其他']
DIRECTIONS = ['寄给顾客', '顾客退回']
KINDS = {'uncollected': '未揽收', 'stalled': '轨迹停滞', 'refused': '拒收', 'damaged': '破损'}
SOURCE_LABELS = {'manual': '人工导入', 'demo': '演示资料'}
HEADERS = {'订单号': 'order_id', '承运商': 'carrier', '运单号': 'tracking_number', '寄送方向': 'direction', '事件时间': 'event_time', '事件类型': 'event_type', '地点': 'location', '事件说明': 'description', '来源事件编号': 'source_event_id', '交寄时间': 'handed_over_at', '资料更新时间': 'checked_until', '资料说明': 'source_note'}
ROW_FIELDS = set(HEADERS.values()) | {'handed_over_basis', 'data_complete', 'event_confirmed', '_row_number', '_identifier_errors'}
MAX_ROWS, MAX_BYTES = 1000, 5 * 1024 * 1024
_json, _hash, _fail, _text = work._json, work._hash, work._fail, work._text


def init_schema(c):
    for statement in SCHEMA.split(';'):
        if statement.strip(): c.execute(statement)
    if not c.execute('SELECT 1 FROM logistics_rules WHERE id=1').fetchone():
        stamp = db.now()
        v = {'version': 1, 'source_label': '本店内部跟进规则', 'confirmed': False, 'basis': '演示建议值，员工核对保存后方可用于人工资料', 'stale_hours': 72, 'updated_at': stamp, 'updated_by': None, 'rules': {k: {'enabled': True, 'hours': h, 'severity': '一般', 'require_event_confirmation': k in ['refused', 'damaged'], 'confirmation_sources': ['manual', 'demo']} for k, h in [('uncollected', 24), ('stalled', 48), ('refused', None), ('damaged', None)]}}
        c.execute('INSERT INTO logistics_rules VALUES (1,?,?)', (1, _json(v)))


def _dt(value): return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(timezone.utc)
def _row(c, table, identity, label):
    if not isinstance(identity, str): _fail(422, f'请选择有效的{label}')
    row = c.execute(f'SELECT data FROM {table} WHERE id=?', (identity,)).fetchone()
    if not row: _fail(404, f'{label}不存在，请刷新后重试')
    return json.loads(row[0])

def _rules(c): return json.loads(c.execute('SELECT data FROM logistics_rules WHERE id=1').fetchone()[0])
def _events(c, identity, effective=False):
    rows = [json.loads(r[0]) for r in c.execute('SELECT data FROM logistics_events WHERE shipment_id=?' + (' AND effective=1' if effective else ''), (identity,))]
    return sorted(rows, key=lambda x: (x['event_time'], x['recorded_at'], x['id']))

def _natural(row): return _hash([row['carrier'].strip(), row['tracking_number'].strip(), row['direction']])
def _content(row): return {k: row.get(k, '') for k in ['event_time', 'event_type', 'location', 'description']}
def _event_keys(row, natural, source_name):
    source = _hash([natural, source_name, row['source_event_id']]) if row.get('source_event_id') else None
    return ('source:' + source if source else 'content:' + _hash([natural, _content(row)])), source


def get_rules(user):
    service.require_staff(user)
    with work._read() as c:
        return {**_rules(c), 'history': [json.loads(r[0]) for r in c.execute('SELECT data FROM logistics_rule_history ORDER BY rowid')]}


def patch_rules(user, data):
    service.require_staff(user); work._keys(data, ['expected_version', 'request_id', 'confirmed', 'basis', 'stale_hours', 'rules', 'reason'])
    with db.transaction() as c:
        rid, digest, old = work._memo(c, user, 'logistics:rules', data)
        if old: return old
        v = _rules(c); before = dict(v)
        if work._version(data.get('expected_version')) != v['version']: _fail(409, '物流规则已更新，请刷新后重新保存')
        if data.get('confirmed') is not True: _fail(422, '请明确确认这些是已核对的本店内部跟进规则')
        basis = _text(data.get('basis'), '本店规则依据说明', 1000); reason = _text(data.get('reason'), '修改原因')
        stale = data.get('stale_hours')
        if type(stale) is not int or not 1 <= stale <= 8760: _fail(422, '资料过旧阈值应为 1–8760 的整数小时')
        rules = data.get('rules'); work._keys(rules, KINDS)
        if set(rules) != set(KINDS): _fail(422, '请完整填写四项物流规则')
        checked = {}
        for kind, rule in rules.items():
            work._keys(rule, ['enabled', 'hours', 'severity', 'require_event_confirmation', 'confirmation_sources'])
            if type(rule.get('enabled')) is not bool or type(rule.get('require_event_confirmation')) is not bool: _fail(422, '请明确选择启停和事件确认要求')
            hours = rule.get('hours')
            if kind in ['uncollected', 'stalled'] and (type(hours) is not int or not 1 <= hours <= 8760): _fail(422, '异常阈值应为 1–8760 的整数小时')
            if kind in ['refused', 'damaged'] and hours is not None: _fail(422, '拒收和破损按明确事件判断，请将小时数留空')
            if rule.get('severity') not in ['一般', '紧急']: _fail(422, '请选择一般或紧急程度')
            sources = rule.get('confirmation_sources', ['manual', 'demo'])
            if not isinstance(sources, list) or any(v not in SOURCE_LABELS for v in sources) or len(sources) != len(set(sources)): _fail(422, '请选择有效且不重复的需要确认的资料来源')
            checked[kind] = dict(rule) | {'confirmation_sources': sources}
        v.update(version=v['version'] + 1, confirmed=True, basis=basis, stale_hours=stale, rules=checked, updated_at=db.now(), updated_by=work._actor(user))
        hist = {'id': db.uid('logistics_rule_history'), 'before': before, 'after': v, 'reason': reason, 'actor': work._actor(user), 'at': db.now()}
        c.execute('UPDATE logistics_rules SET version=?,data=? WHERE id=1', (v['version'], _json(v)))
        c.execute('INSERT INTO logistics_rule_history VALUES (?,?)', (hist['id'], _json(hist)))
        service.audit(c, user, '修改本店物流规则', 'logistics_rules', before, hist)
        return work._remember(c, user, 'logistics:rules', rid, digest, v)


def available_orders(user):
    service.require_staff(user)
    with work._read() as c:
        result = []
        for r in c.execute('SELECT * FROM orders ORDER BY rowid DESC'):
            o = json.loads(r['data']); person = json.loads(c.execute('SELECT data FROM users WHERE id=?', (r['user_id'],)).fetchone()[0])
            result.append({k: o.get(k) for k in ['id', 'product_name', 'sku_label', 'quantity']} | {'user_name': person['name']})
        return {'items': result, 'notice': '订单来自已有本地业务记录；历史模拟签收时间不作为物流轨迹依据。'}


def parse_file(filename, content):
    if len(content) > MAX_BYTES: _fail(422, '文件超过 5 MB，请分成较小文件导入')
    if not content: _fail(422, '文件没有内容，请选择标准表格')
    rows = []
    if filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        try: wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except Exception: _fail(422, '无法读取表格，请保存为标准 .xlsx 文件后重试')
        try:
            sheets = []
            for sheet in wb.worksheets:
                cells = []
                for physical_row, values in enumerate(sheet.iter_rows(), 1):
                    values = list(values)
                    if not any(cell.value is not None and str(cell.value).strip() for cell in values): continue
                    cells.append((physical_row, values))
                    if len(cells) - 1 > MAX_ROWS: _fail(422, '表格超过 1000 行非空资料，请分批导入；本次没有截断或保存')
                if cells: sheets.append(cells)
                if sum(len(s) - 1 for s in sheets) > MAX_ROWS: _fail(422, '表格超过 1000 行非空资料，请分批导入；本次没有截断或保存')
            if len(sheets) != 1: _fail(422, '首版每个文件只支持一个非空工作表，请分别导入')
            cells = sheets[0]
            headers = [str(cell.value).strip() if cell.value is not None else '' for cell in cells[0][1]]
            if any(cell.data_type == 'f' for _, row in cells for cell in row): _fail(422, '表格包含公式，请核对后粘贴为文本或值再导入')
            data_rows = [(row, [cell.value for cell in row], physical_row) for physical_row, row in cells[1:]]
        finally: wb.close()
    elif filename.lower().endswith('.csv'):
        try: decoded = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: decoded = content.decode('gb18030')
            except UnicodeDecodeError: _fail(422, '无法读取文件编码，请保存为 UTF-8 的 CSV')
        try: cells = [(i, r) for i, r in enumerate(csv.reader(io.StringIO(decoded), strict=True), 1) if any(str(v).strip() for v in r)]
        except csv.Error: _fail(422, 'CSV 格式不完整，请按模板重新保存')
        if not cells: _fail(422, '表格没有内容')
        headers = [s.strip() for s in cells[0][1]]; data_rows = [(None, r, i) for i, r in cells[1:]]
    else: _fail(422, '仅支持 .xlsx 和 .csv 标准表格')
    if len(data_rows) > MAX_ROWS: _fail(422, '表格超过 1000 行非空资料，请分批导入；本次没有截断或保存')
    if len(set(h for h in headers if h)) != len([h for h in headers if h]): _fail(422, '表格列名重复，请按中文模板核对')
    if any(h and h not in HEADERS for h in headers): _fail(422, '存在不支持的列名，请使用中文模板；额外列不会被静默丢弃')
    if any(h not in headers for h in ['订单号', '承运商', '运单号', '寄送方向', '事件时间', '事件类型', '资料更新时间']): _fail(422, '缺少必要列，请下载中文模板后填写')
    for raw, values, index in data_rows:
        if len(values) > len(headers) and any(v is not None and str(v).strip() for v in values[len(headers):]): _fail(422, f'第 {index} 行超出模板列范围，请核对')
        row = {'_row_number': index, '_identifier_errors': []}
        for j, value in enumerate(values[:len(headers)]):
            if not headers[j]:
                if value is not None and str(value).strip(): _fail(422, f'第 {index} 行存在无标题的非空列，请补充中文标题')
                continue
            field = HEADERS[headers[j]]
            if field in ['order_id', 'tracking_number', 'source_event_id'] and value is not None and not isinstance(value, str): row['_identifier_errors'].append(headers[j])
            if isinstance(value, datetime): value = value.isoformat()
            row[field] = '' if value is None else value
        rows.append(row)
    return rows, hashlib.sha256(content).hexdigest()


def _parse_time(value, label, issues, naive_ok, required=False):
    if value is None or value == '':
        if required: issues.append(label + '缺失，请填写资料实际时间')
        return None, None
    if not isinstance(value, str): issues.append(label + '格式无效，请填写时间'); return None, None
    try:
        value = value.strip(); dt = datetime.fromisoformat(value.replace('Z', '+00:00')); original_zone = str(dt.tzinfo) if dt.tzinfo else 'Asia/Shanghai（预览明确采用）'
        if dt.tzinfo is None:
            if not naive_ok: issues.append(label + '没有时区，请在预览明确采用中国时区'); return None, original_zone
            dt = dt.replace(tzinfo=ZoneInfo('Asia/Shanghai'))
        dt = dt.astimezone(timezone.utc)
        if dt > _dt(db.now()): issues.append(label + '在未来，请核对原始资料')
        return dt.isoformat(), original_zone
    except (ValueError, OverflowError): issues.append(label + '无法识别，请使用年-月-日 时:分'); return None, None


def _dependencies(c, rows):
    orders, shipments = {}, {}
    for row in rows:
        if row.get('order_id'):
            r = c.execute('SELECT * FROM orders WHERE id=?', (row['order_id'],)).fetchone(); orders[row['order_id']] = dict(r) if r else None
        if all(row.get(k) for k in ['carrier', 'tracking_number', 'direction']):
            key = _natural(row); s = c.execute('SELECT * FROM logistics_shipments WHERE natural_key=?', (key,)).fetchone()
            shipments[key] = {'shipment': dict(s), 'events': [dict(r) for r in c.execute('SELECT * FROM logistics_events WHERE shipment_id=? ORDER BY id', (s['id'],))]} if s else None
    return _hash({'orders': orders, 'shipments': shipments, 'rules': _rules(c)})


def import_preview(user, data, file_fingerprint=None):
    service.require_staff(user); work._keys(data, ['rows', 'source_kind', 'source_name', 'source_note', 'data_complete', 'event_confirmed', 'confirm_china_timezone', 'request_id'])
    rows = data.get('rows')
    if not isinstance(rows, list) or not rows or len(rows) > MAX_ROWS: _fail(422, '请填写 1–1000 行物流资料')
    if len(_json(data).encode()) > MAX_BYTES: _fail(422, '资料超过 5 MB，请分批导入')
    source_kind = data.get('source_kind', 'manual')
    if source_kind not in SOURCE_LABELS: _fail(422, '请选择人工导入或演示资料')
    source_name = _text(data.get('source_name'), '资料来源名称', 200)
    note = _text(data.get('source_note', ''), '资料来源说明', 1000, False)
    for flag in ['data_complete', 'event_confirmed', 'confirm_china_timezone']:
        if flag in data and type(data[flag]) is not bool: _fail(422, '请明确选择资料核对选项')
    with db.transaction() as c:
        rid, digest, old = work._memo(c, user, 'logistics:preview', data | {'file_fingerprint': file_fingerprint})
        if old: return old
        results, normalized, seen = [], [], {}
        for number, raw in enumerate(rows, 1):
            if not isinstance(raw, dict): _fail(422, f'第 {number} 行不是有效表格资料')
            raw = {HEADERS.get(k, k): v for k, v in raw.items()}; work._keys(raw, ROW_FIELDS)
            idx = raw.get('_row_number', number); issues, warnings = [], []; row = {}
            for field in ['order_id', 'carrier', 'tracking_number', 'direction', 'event_type', 'location', 'description', 'source_event_id', 'source_note', 'handed_over_basis']:
                val = raw.get(field, '')
                if val is None: val = ''
                if not isinstance(val, str): issues.append(next((k for k, v in HEADERS.items() if v == field), field) + '必须保存为文本，保留前导零；请将整列设置为文本并重新填写'); val = ''
                if len(val) > (2000 if field in ['description', 'source_note', 'handed_over_basis'] else 200): issues.append('字段过长，请缩短原始资料说明后重试')
                row[field] = val.strip()
            for field in raw.get('_identifier_errors', []): issues.append(field + '曾作为数值保存，不能恢复前导零；请设置文本格式后按原凭证重新填写')
            if not all(row[k] for k in ['order_id', 'carrier', 'tracking_number']): issues.append('订单号、承运商、运单号不能为空')
            if row['direction'] not in DIRECTIONS: issues.append('寄送方向请选择寄给顾客或顾客退回')
            if row['event_type'] not in EVENT_TYPES: issues.append('事件类型请选择明确的中文类型；自由说明不会自动识别拒收或破损')
            row['event_time'], row['original_timezone'] = _parse_time(raw.get('event_time'), '事件时间', issues, data.get('confirm_china_timezone') is True, True)
            row['checked_until'], row['checked_original_timezone'] = _parse_time(raw.get('checked_until'), '资料更新时间', issues, data.get('confirm_china_timezone') is True, True)
            row['handed_over_at'], _ = _parse_time(raw.get('handed_over_at'), '交寄时间', issues, data.get('confirm_china_timezone') is True)
            row['handed_over_basis'] = row['handed_over_basis'] or row['source_note'] or note
            if row['handed_over_at'] and not row['handed_over_basis']: issues.append('有交寄时间时请填写交寄依据或资料说明')
            if row['checked_until'] and row['event_time'] and _dt(row['event_time']) > _dt(row['checked_until']): issues.append('事件时间晚于资料核对时间，请核对')
            if row['handed_over_at'] and row['checked_until'] and _dt(row['handed_over_at']) > _dt(row['checked_until']): issues.append('交寄时间晚于资料核对时间，请核对')
            for flag in ['data_complete', 'event_confirmed']:
                row[flag] = raw.get(flag, data.get(flag, False))
                if type(row[flag]) is not bool: issues.append('资料完整和事件确认需明确选择'); row[flag] = False
            if not row['data_complete']: warnings.append('资料未确认完整，时间规则将保留为资料不足')
            if _uncertain_event(row): warnings.append('原说明未明确证实此事件，轨迹与异常判断将显示待核对')
            if not row['handed_over_at']: warnings.append('未登记有依据的交寄时间，不能判断未揽收')
            order = c.execute('SELECT * FROM orders WHERE id=?', (row['order_id'],)).fetchone()
            if not order: issues.append('订单不存在，请从已有订单选择，不能按姓名推测')
            key, source = (None, None)
            existing = None
            if not issues:
                natural = _natural(row); shipment = c.execute('SELECT * FROM logistics_shipments WHERE natural_key=?', (natural,)).fetchone()
                if shipment and shipment['order_id'] != row['order_id']: issues.append('该运单已关联另一订单，请核对')
                key, source = _event_keys(row, natural, source_name)
                existing = c.execute('SELECT data FROM logistics_events WHERE dedupe_key=? OR source_event_key=?', (key, source)).fetchone()
                earlier = seen.get(key)
                if existing and _content(json.loads(existing[0])) != _content(row): issues.append('相同来源事件编号已有不同内容，请使用历史更正入口，不能覆盖')
                if earlier and (_content(earlier) != _content(row) or earlier['order_id'] != row['order_id']): issues.append('本批相同运单或来源事件编号出现不同内容，请核对')
                if shipment:
                    old_shipment = json.loads(shipment['data'])
                    if old_shipment['source_kind'] != source_kind: issues.append('同一运单的演示与人工资料类别冲突，请核对来源')
                    if row['handed_over_at'] and old_shipment.get('handed_over_at') and row['handed_over_at'] != old_shipment['handed_over_at']: issues.append('交寄时间与已登记依据冲突，请核对后使用运单资料更正')
                    same = [e for e in _events(c, shipment['id'], True) if e['event_time'] == row['event_time'] and e['event_type'] != row['event_type'] and e['event_type'] not in ['破损', '其他'] and row['event_type'] not in ['破损', '其他']]
                    if same: warnings.append('同一时刻有不同状态，轨迹将显示待核对，不以最后一行判定')
                for e in normalized:
                    if _natural(e) == natural and e['event_time'] == row['event_time'] and e['event_type'] != row['event_type'] and e['event_type'] not in ['破损', '其他'] and row['event_type'] not in ['破损', '其他']: warnings.append('本批同一时刻有不同状态，轨迹将显示待核对')
            status = 'invalid' if issues else 'duplicate' if existing or key in seen else 'new'
            if status == 'invalid' and any('冲突' in s or '另一订单' in s or '不同内容' in s for s in issues): status = 'conflict'
            row['_row_number'] = idx; row['dedupe_key'] = key; row['source_event_key'] = source
            results.append({'row_number': idx, 'status': status, 'status_label': {'new': '新增', 'duplicate': '重复，将跳过', 'conflict': '冲突，须修正', 'invalid': '无效，须修正'}[status], 'issues': issues, 'warnings': list(dict.fromkeys(warnings)), 'row': row})
            if not issues: normalized.append(row); seen[key] = row
        counts = {k: sum(r['status'] == k for r in results) for k in ['new', 'duplicate', 'conflict', 'invalid']}
        rules = _rules(c); expires = (_dt(db.now()) + timedelta(minutes=15)).isoformat()
        result = {'preview_id': db.uid('tracking_preview'), 'fingerprint': _hash({'rows': rows, 'source': source_kind, 'source_name': source_name, 'rules': rules['version']}), 'rule_version': rules['version'], 'rows': results, 'counts': counts, 'can_confirm': not counts['invalid'] and not counts['conflict'], 'expires_at': expires, 'timezone': 'Asia/Shanghai', 'source_label': SOURCE_LABELS[source_kind], 'notice': '预览不保存正式轨迹；确认只登记本地资料和内部异常，不修改订单签收、库存或售后。'}
        stored = result | {'normalized_rows': normalized, 'source_kind': source_kind, 'source_name': source_name, 'source_note': note, 'file_fingerprint': file_fingerprint or _hash(rows), 'dependencies': _dependencies(c, normalized)}
        c.execute('INSERT INTO logistics_previews VALUES (?,?,?,?,?)', (result['preview_id'], user['id'], expires, 0, _json(stored)))
        return work._remember(c, user, 'logistics:preview', rid, digest, result)


def _uncertain_event(event):
    marker = {'拒收': '拒收', '破损': '破损'}.get(event.get('event_type'))
    if not marker: return False
    text = event.get('description', '')
    return bool(re.search(r'(?:可能|疑似|好像|不确定|是否|尚未确认|待核实).{0,8}' + marker + '|' + marker + r'.{0,8}(?:待核实|不确定|可能|待确认)|(?:没有|并未|未见|未发现|无)' + marker, text))


def _ambiguities(events):
    times = {}
    for event in events: times.setdefault(event['event_time'], set()).add(event['event_type'])
    return {at for at, types in times.items() if len(types - {'破损', '其他'}) > 1}


def _shipment_view(c, shipment):
    events = _events(c, shipment['id'], True); ambiguous = _ambiguities(events)
    latest = events[-1] if events else None
    terminal = [e for e in events if e['event_type'] in ['已签收', '退回'] and e['event_time'] not in ambiguous and e.get('event_confirmed')]
    # Old tracking progress cannot undo a later terminal event.
    status = '待核对' if latest and (latest['event_time'] in ambiguous or _uncertain_event(latest) or latest['event_type'] in ['已签收', '拒收', '破损', '退回'] and not latest.get('event_confirmed')) else terminal[-1]['event_type'] if terminal else latest['event_type'] if latest else '资料不足'
    stale = not shipment.get('checked_until') or (_dt(db.now()) - _dt(shipment['checked_until'])).total_seconds() > _rules(c)['stale_hours'] * 3600
    return shipment | {'latest_event_at': latest['event_time'] if latest else None, 'tracking_status': status, 'ambiguous_times': sorted(ambiguous), 'data_stale': stale, 'source_label': SOURCE_LABELS[shipment['source_kind']], 'events': events}


def _judgement(c, shipment, rules, kind):
    events = _events(c, shipment['id'], True); ambiguous = _ambiguities(events); rule = rules['rules'][kind]
    checked = shipment.get('checked_until'); event = None; basis_at = None; outcome, reason = 'normal', ''
    if not rule['enabled']: outcome, reason = 'disabled', '本店已停用此规则'
    elif shipment['source_kind'] != 'demo' and not rules['confirmed']: outcome, reason = 'insufficient', '本店规则尚未由员工核对确认'
    elif not checked or kind in ['uncollected', 'stalled'] and not shipment.get('data_complete'): outcome, reason = 'insufficient', '资料更新时间或完整性尚未核对'
    elif kind in ['uncollected', 'stalled']:
        valid = [e for e in events if e['event_time'] not in ambiguous]
        if ambiguous and (not valid or max(ambiguous) >= valid[-1]['event_time']): outcome, reason = 'insufficient', '最近同刻事件相互矛盾，待核对'
        else:
            collected = any(e['event_type'] in ['已揽收', '运输中', '派送中', '已签收', '退回'] for e in valid)
            terminal = any(e['event_type'] in ['已签收', '退回'] and e.get('event_confirmed') for e in valid)
            if any(e['event_type'] in ['已签收', '退回'] and not e.get('event_confirmed') for e in valid):
                outcome, reason = 'insufficient', '签收或退回类型尚未核对，不能认定运输已结束'
                valid = []
            if outcome == 'insufficient': pass
            elif kind == 'uncollected':
                basis_at = shipment.get('handed_over_at')
                if collected: reason = '已有有效揽收或运输证据'
                elif not basis_at or not shipment.get('handed_over_basis'): outcome, reason = 'insufficient', '缺少有依据的交寄时间，不以订单时间补算'
                elif (_dt(checked) - _dt(basis_at)).total_seconds() > rule['hours'] * 3600: outcome, reason = 'triggered', '截至资料核对时间仍未见揽收，已超过本店阈值'
                else: reason = '截至资料核对时间尚未超过本店阈值'
            else:
                progress = [e for e in valid if e['event_type'] in ['已揽收', '运输中', '派送中']]
                if terminal: reason = '已有签收或退回结束证据；争议不自动关闭'
                elif not collected or not progress: outcome, reason = 'insufficient', '缺少有效揽收或运输进展时间'
                else:
                    event = progress[-1]; basis_at = event['event_time']
                    outcome = 'triggered' if (_dt(checked) - _dt(basis_at)).total_seconds() > rule['hours'] * 3600 else 'normal'
                    reason = '截至资料核对时间运输进展已超过本店阈值' if outcome == 'triggered' else '截至资料核对时间尚未超过本店阈值'
    else:
        matches = [e for e in events if e['event_type'] == {'refused': '拒收', 'damaged': '破损'}[kind]]
        if matches:
            event = matches[-1]; basis_at = event['event_time']
            if event['event_time'] in ambiguous: outcome, reason = 'insufficient', '同刻状态冲突，先核对事件，不认定责任'
            elif _uncertain_event(event): outcome, reason = 'insufficient', '原说明仍为可能、疑似或否定事件，须先核实来源，不能认定责任'
            elif rule['require_event_confirmation'] and event['source_kind'] in rule.get('confirmation_sources', ['manual', 'demo']) and not event.get('event_confirmed'): outcome, reason = 'insufficient', '事件来源尚未明确核对确认'
            else: outcome, reason = 'triggered', '存在已核对的明确事件，不自动认定责任或执行售后'
        else: reason = '未见明确事件；不会从自由说明推定'
    return {'id': db.uid('logistics_judgement'), 'shipment_id': shipment['id'], 'kind': kind, 'kind_label': KINDS[kind], 'outcome': outcome, 'reason': reason, 'basis_at': basis_at, 'trigger_event_id': event['id'] if event else None, 'checked_until': checked, 'judged_at': db.now(), 'latest_event_at': events[-1]['event_time'] if events else None, 'data_stale': not checked or (_dt(db.now()) - _dt(checked)).total_seconds() > rules['stale_hours'] * 3600, 'rule_version': rules['version'], 'rule_snapshot': {'source_label': rules['source_label'], 'basis': rules['basis'], 'confirmed': rules['confirmed'], 'stale_hours': rules['stale_hours'], **rule}, 'source_kind': shipment['source_kind'], 'import_ids': sorted({e['import_id'] for e in events}), 'evidence_fingerprint': _hash({'events': [{k:e.get(k) for k in ['id', 'event_time', 'event_type', 'event_confirmed']} for e in events], 'checked_until': checked, 'handed_over_at': shipment.get('handed_over_at'), 'handed_over_basis': shipment.get('handed_over_basis'), 'data_complete': shipment.get('data_complete')})}


def _recheck(c, user, shipment, source='rule'):
    rules = _rules(c); findings = []; exceptions = []
    for kind in KINDS:
        finding = _judgement(c, shipment, rules, kind); finding['actor'] = work._actor(user); finding['source'] = source
        active = c.execute("SELECT data FROM logistics_exceptions WHERE shipment_id=? AND kind=? AND status!='closed'", (shipment['id'], kind)).fetchone()
        value = json.loads(active[0]) if active else None
        if finding['outcome'] == 'triggered' and value is None:
            latest_closed = c.execute("SELECT data FROM logistics_exceptions WHERE shipment_id=? AND kind=? AND status='closed' ORDER BY rowid DESC LIMIT 1", (shipment['id'], kind)).fetchone()
            if latest_closed and json.loads(latest_closed[0])['latest_judgement'].get('evidence_fingerprint') == finding['evidence_fingerprint']:
                finding.update(outcome='previously_closed', reason='同一份证据对应异常已关闭；新增有效资料后才建立新的跟进轮次')
        if finding['outcome'] == 'triggered':
            if value:
                value = value | {'version': value['version'] + 1, 'last_found_at': finding['judged_at'], 'latest_judgement': finding, 'updated_at': db.now()}
                c.execute('UPDATE logistics_exceptions SET version=?,data=? WHERE id=?', (value['version'], _json(value), value['id']))
            else:
                previous = c.execute('SELECT id FROM logistics_exceptions WHERE shipment_id=? AND kind=? ORDER BY rowid DESC LIMIT 1', (shipment['id'], kind)).fetchone()
                value = {'id': db.uid('logistics_exception'), 'shipment_id': shipment['id'], 'order_id': shipment['order_id'], 'user_id': shipment['user_id'], 'kind': kind, 'kind_label': KINDS[kind], 'status': 'open', 'version': 1, 'source': 'demo' if shipment['source_kind'] == 'demo' else source, 'source_label': '演示资料' if shipment['source_kind'] == 'demo' else '员工发现' if source == 'staff' else '按本店规则发现', 'first_found_at': finding['judged_at'], 'last_found_at': finding['judged_at'], 'updated_at': db.now(), 'original_judgement': finding, 'latest_judgement': finding, 'previous_exception_id': previous[0] if previous else None, 'linked_target': None, 'closure': None}
                c.execute('INSERT INTO logistics_exceptions VALUES (?,?,?,?,?,?)', (value['id'], shipment['id'], kind, 'open', 1, _json(value)))
            exceptions.append(value); finding['exception_id'] = value['id']
        else:
            finding['exception_id'] = value['id'] if value else None
            if value:
                value.update(version=value['version'] + 1, latest_judgement=finding, updated_at=db.now())
                c.execute('UPDATE logistics_exceptions SET version=?,data=? WHERE id=?', (value['version'], _json(value), value['id']))
        c.execute('INSERT INTO logistics_judgements VALUES (?,?,?,?)', (finding['id'], shipment['id'], finding['exception_id'], _json(finding)))
        findings.append(finding)
    return {'judgements': findings, 'exceptions': exceptions}


def import_confirm(user, data):
    service.require_staff(user); work._keys(data, ['preview_id', 'fingerprint', 'rule_version', 'request_id'])
    with db.transaction() as c:
        rid, digest, old = work._memo(c, user, 'logistics:confirm', data)
        if old: return old
        row = c.execute('SELECT * FROM logistics_previews WHERE id=? AND actor_id=?', (data.get('preview_id'), user['id'])).fetchone()
        if not row: _fail(404, '预览不存在或不属于当前员工，请重新预览')
        v = json.loads(row['data'])
        if row['consumed']: _fail(409, '这份预览已确认，请查看原回执')
        if _dt(row['expires_at']) <= _dt(db.now()): _fail(409, '预览已过期，请重新预览')
        if not v['can_confirm']: _fail(422, '请先修正所有无效和冲突行，再重新预览')
        if data.get('fingerprint') != v['fingerprint'] or data.get('rule_version') != v['rule_version'] or _dependencies(c, v['normalized_rows']) != v['dependencies']: _fail(409, '物流资料、订单或规则已更新，请重新预览后确认')
        stamp = db.now(); batch = {'id': db.uid('tracking_import'), 'source_kind': v['source_kind'], 'source_label': SOURCE_LABELS[v['source_kind']], 'source_name': v['source_name'], 'source_note': v['source_note'], 'file_fingerprint': v['file_fingerprint'], 'preview_fingerprint': v['fingerprint'], 'rule_version': v['rule_version'], 'actor': work._actor(user), 'recorded_at': stamp, 'counts': v['counts'], 'checked_until': max(r['checked_until'] for r in v['normalized_rows']), 'request_id': rid}
        shipments, inserted, duplicates = {}, [], []
        for r in v['normalized_rows']:
            natural = _natural(r); current = c.execute('SELECT data FROM logistics_shipments WHERE natural_key=?', (natural,)).fetchone()
            if current:
                shipment = json.loads(current[0])
                if shipment['order_id'] != r['order_id']: _fail(409, '该运单已关联另一订单，请核对')
                if shipment['source_kind'] != v['source_kind']: _fail(409, '该运单已有不同资料类别，演示与人工资料不能混写')
            else:
                o = c.execute('SELECT * FROM orders WHERE id=?', (r['order_id'],)).fetchone()
                shipment = {'id': db.uid('shipment'), 'natural_key': natural, 'carrier': r['carrier'], 'tracking_number': r['tracking_number'], 'direction': r['direction'], 'order_id': r['order_id'], 'user_id': o['user_id'], 'source_kind': v['source_kind'], 'handed_over_at': r['handed_over_at'], 'handed_over_basis': r['handed_over_basis'], 'checked_until': r['checked_until'], 'data_complete': r['data_complete'], 'recorded_by': work._actor(user), 'created_at': stamp, 'updated_at': stamp, 'version': 0, 'import_ids': [], 'history': []}
            before = {k: shipment.get(k) for k in ['handed_over_at', 'handed_over_basis', 'checked_until', 'data_complete']}
            if not shipment.get('handed_over_at') and r['handed_over_at']: shipment.update(handed_over_at=r['handed_over_at'], handed_over_basis=r['handed_over_basis'])
            if _dt(r['checked_until']) >= _dt(shipment['checked_until']): shipment.update(checked_until=r['checked_until'], data_complete=r['data_complete'])
            shipment.update(version=shipment['version'] + 1, updated_at=stamp, import_ids=list(dict.fromkeys([*shipment['import_ids'], batch['id']])))
            shipment['history'] = [*shipment.get('history', []), {'at': stamp, 'actor': work._actor(user), 'action': '导入资料', 'import_id': batch['id'], 'before': before, 'after': {k: shipment.get(k) for k in before}}]
            c.execute('INSERT INTO logistics_shipments VALUES (?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET version=excluded.version,data=excluded.data', (shipment['id'], natural, shipment['order_id'], shipment['user_id'], shipment['source_kind'], shipment['version'], _json(shipment)))
            exists = c.execute('SELECT id FROM logistics_events WHERE dedupe_key=? OR source_event_key=?', (r['dedupe_key'], r['source_event_key'])).fetchone()
            if exists: duplicates.append({'row_number': r['_row_number'], 'event_id': exists[0]})
            else:
                event = {k: r.get(k) for k in ['event_time', 'event_type', 'location', 'description', 'source_event_id', 'original_timezone', 'event_confirmed']}
                event.update(id=db.uid('shipment_event'), shipment_id=shipment['id'], source_name=v['source_name'], source_kind=v['source_kind'], import_id=batch['id'], source_row=r['_row_number'], recorded_at=stamp, recorded_by=work._actor(user), effective=True, replaced_by=None, replaces=None, version=1, correction_reason=None)
                c.execute('INSERT INTO logistics_events VALUES (?,?,?,?,?,?)', (event['id'], shipment['id'], r['dedupe_key'], r['source_event_key'], 1, _json(event)))
                inserted.append({'row_number': r['_row_number'], 'event_id': event['id']})
            shipments[shipment['id']] = shipment
        batch.update(inserted=inserted, duplicates=duplicates, shipment_ids=list(shipments), counts={'new': len(inserted), 'duplicate': len(duplicates), 'conflict': 0, 'invalid': 0})
        c.execute('INSERT INTO logistics_imports VALUES (?,?,?,?)', (batch['id'], user['id'], batch['file_fingerprint'], _json(batch)))
        findings = [_recheck(c, user, shipment) for shipment in shipments.values()]
        result = {'import': batch, 'shipments': [_shipment_view(c, s) for s in shipments.values()], 'judgements': [j for f in findings for j in f['judgements']], 'exceptions': [e for f in findings for e in f['exceptions']], 'notice': '已保存本地物流资料及内部跟进，未发顾客消息，未执行退款、补发或出库。'}
        c.execute('UPDATE logistics_previews SET consumed=1 WHERE id=?', (v['preview_id'],))
        service.audit(c, user, '确认物流资料导入', batch['id'], None, batch)
        return work._remember(c, user, 'logistics:confirm', rid, digest, result)


def list_shipments(user, q='', order_id=None, source_kind='all', exception_only=False):
    service.require_staff(user)
    if source_kind not in ['all', *SOURCE_LABELS]: _fail(422, '请选择有效资料来源')
    with work._read() as c:
        result = []
        for row in c.execute('SELECT data FROM logistics_shipments ORDER BY rowid DESC'):
            s = json.loads(row[0])
            if order_id and s['order_id'] != order_id: continue
            if source_kind != 'all' and s['source_kind'] != source_kind: continue
            if q and q.strip().casefold() not in (' '.join([s['carrier'], s['tracking_number'], s['order_id']])).casefold(): continue
            ex = [json.loads(r[0]) for r in c.execute("SELECT data FROM logistics_exceptions WHERE shipment_id=? AND status!='closed'", (s['id'],))]
            if exception_only and not ex: continue
            view = _shipment_view(c, s); view.pop('events'); view['exceptions'] = ex; result.append(view)
        return {'items': result, 'total': len(result), 'as_of': db.now(), 'timezone': 'Asia/Shanghai', 'notice': '本地登记资料；资料更新时间是核对到的时间，页面时间不代表承运商实时更新。'}


def get_shipment(user, identity):
    service.require_staff(user)
    with work._read() as c:
        s = _row(c, 'logistics_shipments', identity, '运单'); ex = [json.loads(r[0]) for r in c.execute('SELECT data FROM logistics_exceptions WHERE shipment_id=? ORDER BY rowid DESC', (identity,))]
        return _shipment_view(c, s) | {'event_history': _events(c, identity), 'imports': [_row(c, 'logistics_imports', i, '导入批次') for i in s['import_ids']], 'exceptions': ex, 'judgements': [json.loads(r[0]) for r in c.execute('SELECT data FROM logistics_judgements WHERE shipment_id=? ORDER BY rowid DESC', (identity,))]}


def recheck(user, data):
    service.require_staff(user); work._keys(data, ['shipment_ids', 'request_id', 'reason', 'source'])
    ids = data.get('shipment_ids')
    if not isinstance(ids, list) or not 1 <= len(ids) <= MAX_ROWS or any(not isinstance(i, str) for i in ids): _fail(422, '请选择 1–1000 票现有运单')
    reason = _text(data.get('reason'), '重新核对原因'); source = data.get('source', 'rule')
    if source not in ['rule', 'staff']: _fail(422, '请选择员工发现或按本店规则发现')
    with db.transaction() as c:
        rid, digest, old = work._memo(c, user, 'logistics:recheck', data)
        if old: return old
        findings = [_recheck(c, user, _row(c, 'logistics_shipments', identity, '运单'), source) for identity in dict.fromkeys(ids)]
        result = {'judgements': [j for f in findings for j in f['judgements']], 'exceptions': [e for f in findings for e in f['exceptions']], 'reason': reason, 'as_of': db.now()}
        service.audit(c, user, '重新判断本地物流资料', ','.join(ids), None, result)
        return work._remember(c, user, 'logistics:recheck', rid, digest, result)


def exception_anchor(c, identity):
    value = _row(c, 'logistics_exceptions', identity, '物流异常'); shipment = _row(c, 'logistics_shipments', value['shipment_id'], '运单')
    person = _row(c, 'users', value['user_id'], '顾客')
    return {'id': value['id'], 'type': 'logistics_exception', 'stage': 'closed' if value['status'] == 'closed' else 'pending', 'status': value['status'], 'conversation_id': None, 'user_id': value['user_id'], 'user_name': person['name'], 'order_id': value['order_id'], 'request_type': value['kind'], 'title': '物流异常 · ' + value['kind_label'], 'reason': value['latest_judgement']['reason'], 'updated_at': value['updated_at'], 'source_label': value['source_label'], 'source_kind': shipment['source_kind'], 'shipment_id': shipment['id'], 'tracking_number': shipment['tracking_number'], 'intent_ids': [], 'linked_target': value.get('linked_target')}


def queue_items(c):
    if not c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='logistics_exceptions'").fetchone(): return []
    return [exception_anchor(c, r[0]) for r in c.execute("SELECT id FROM logistics_exceptions WHERE status!='closed' ORDER BY rowid DESC")]


def list_exceptions(user, status='open', source_kind='all', q=''):
    service.require_staff(user)
    if status not in ['all', 'open', 'closed'] or source_kind not in ['all', 'manual', 'demo']: _fail(422, '请选择有效的物流异常筛选')
    with work._read() as c:
        result = []
        for row in c.execute('SELECT data FROM logistics_exceptions ORDER BY rowid DESC'):
            v = json.loads(row[0]); s = _row(c, 'logistics_shipments', v['shipment_id'], '运单')
            if status != 'all' and v['status'] != status: continue
            if source_kind != 'all' and s['source_kind'] != source_kind: continue
            if q and q.strip().casefold() not in (' '.join([s['carrier'], s['tracking_number'], s['order_id'], v['kind_label']])).casefold(): continue
            item = work._enrich(c, exception_anchor(c, v['id']))
            result.append(v | {'shipment': _shipment_view(c, s), 'item': item, 'history': [json.loads(r[0]) for r in c.execute('SELECT data FROM logistics_exception_history WHERE exception_id=? ORDER BY rowid', (v['id'],))]})
        return {'items': result, 'total': len(result), 'as_of': db.now()}


def exception_action(user, identity, data):
    service.require_staff(user); work._keys(data, ['request_id', 'expected_version', 'action', 'reason', 'evidence', 'target'])
    with db.transaction() as c:
        rid, digest, old = work._memo(c, user, 'logistics:exception:' + identity, data)
        if old: return old
        v = _row(c, 'logistics_exceptions', identity, '物流异常'); before = dict(v)
        if work._version(data.get('expected_version')) != v['version']: _fail(409, '异常资料已更新，请刷新后重试')
        if v['status'] == 'closed': _fail(409, '异常已关闭，请查看历史；新证据可重新判断建立新轮次')
        action = data.get('action'); reason = _text(data.get('reason'), '核对原因'); evidence = _text(data.get('evidence'), '核对证据', 2000)
        if action == 'link':
            item = work._resolve(c, user, data.get('target'))
            if item['type'] == 'logistics_exception' or item['user_id'] != v['user_id'] or item.get('order_id') != v['order_id']: _fail(422, '关联事项必须属于同一顾客及同一订单，请核对具体事项')
            if not work._open(item): _fail(409, '关联事项已结束，请选择仍在跟进的具体事项')
            v['linked_target'] = {'type': item['type'], 'id': item['id'], 'title': item['title'], 'reason': reason, 'evidence': evidence, 'linked_at': db.now(), 'actor': work._actor(user)}
        elif action == 'unlink':
            if not v.get('linked_target'): _fail(409, '尚无关联事项')
            v['linked_target'] = None
        elif action in ['recovered', 'false_alarm', 'transferred']:
            v.update(status='closed', closure={'kind': action, 'label': {'recovered': '已核实恢复', 'false_alarm': '误报', 'transferred': '转其他事项'}[action], 'reason': reason, 'evidence': evidence, 'actor': work._actor(user), 'at': db.now()})
        else: _fail(422, '请选择关联、取消关联或明确的关闭原因')
        v.update(version=v['version'] + 1, updated_at=db.now()); hist = {'id': db.uid('logistics_history'), 'action': action, 'reason': reason, 'evidence': evidence, 'actor': work._actor(user), 'at': db.now(), 'before': before, 'after': v}
        c.execute('UPDATE logistics_exceptions SET status=?,version=?,data=? WHERE id=?', (v['status'], v['version'], _json(v), identity))
        c.execute('INSERT INTO logistics_exception_history VALUES (?,?,?)', (hist['id'], identity, _json(hist)))
        service.audit(c, user, '记录物流异常跟进', identity, before, v)
        return work._remember(c, user, 'logistics:exception:' + identity, rid, digest, v)


def correct_event(user, identity, data):
    service.require_staff(user); work._keys(data, ['request_id', 'expected_version', 'reason', 'replacement', 'confirm_china_timezone'])
    with db.transaction() as c:
        rid, digest, old = work._memo(c, user, 'logistics:event:' + identity, data)
        if old: return old
        event = _row(c, 'logistics_events', identity, '轨迹事件'); before = dict(event)
        if work._version(data.get('expected_version')) != event['version'] or not event['effective']: _fail(409, '事件已更正，请刷新历史后重试')
        reason = _text(data.get('reason'), '更正原因'); replacement = data.get('replacement'); new = None
        if replacement is not None:
            work._keys(replacement, ['event_time', 'event_type', 'location', 'description', 'event_confirmed']); issues = []
            at, original_zone = _parse_time(replacement.get('event_time'), '事件时间', issues, data.get('confirm_china_timezone') is True, True)
            if issues: _fail(422, '；'.join(issues))
            if replacement.get('event_type') not in EVENT_TYPES or type(replacement.get('event_confirmed')) is not bool: _fail(422, '请明确选择事件类型及核对确认')
            s = _row(c, 'logistics_shipments', event['shipment_id'], '运单')
            if at > s['checked_until']: _fail(422, '替代事件晚于资料核对时间，请先导入更新资料')
            new = event | {'id': db.uid('shipment_event'), 'event_time': at, 'original_timezone': original_zone, 'event_type': replacement['event_type'], 'event_confirmed': replacement['event_confirmed'], 'location': _text(replacement.get('location', ''), '地点', 200, False), 'description': _text(replacement.get('description', ''), '原始说明', 2000, False), 'source_event_id': None, 'recorded_at': db.now(), 'recorded_by': work._actor(user), 'replaces': identity, 'replaced_by': None, 'version': 1, 'correction_reason': reason}
            c.execute('INSERT INTO logistics_events VALUES (?,?,?,?,?,?)', (new['id'], event['shipment_id'], 'correction:' + new['id'], None, 1, _json(new)))
        event.update(effective=False, replaced_by=new['id'] if new else None, correction_reason=reason, corrected_at=db.now(), corrected_by=work._actor(user), version=event['version'] + 1)
        c.execute('UPDATE logistics_events SET effective=0,data=? WHERE id=?', (_json(event), identity))
        shipment = _row(c, 'logistics_shipments', event['shipment_id'], '运单'); shipment.update(version=shipment['version'] + 1, updated_at=db.now())
        c.execute('UPDATE logistics_shipments SET version=?,data=? WHERE id=?', (shipment['version'], _json(shipment), shipment['id']))
        judgement = _recheck(c, user, shipment, 'staff'); result = {'original': event, 'replacement': new, **judgement}
        service.audit(c, user, '更正轨迹事件并保留原证据', identity, before, result)
        return work._remember(c, user, 'logistics:event:' + identity, rid, digest, result)


def patch_shipment(user, identity, data):
    service.require_staff(user); work._keys(data, ['request_id', 'expected_version', 'reason', 'handed_over_at', 'handed_over_basis', 'checked_until', 'data_complete', 'confirm_china_timezone'])
    with db.transaction() as c:
        rid, digest, old = work._memo(c, user, 'logistics:shipment:' + identity, data)
        if old: return old
        s = _row(c, 'logistics_shipments', identity, '运单'); before = dict(s)
        if work._version(data.get('expected_version')) != s['version']: _fail(409, '运单资料已更新，请刷新后重试')
        reason = _text(data.get('reason'), '更正原因'); issues = []
        at, _ = _parse_time(data.get('handed_over_at'), '交寄时间', issues, data.get('confirm_china_timezone') is True)
        checked, _ = _parse_time(data.get('checked_until'), '资料更新时间', issues, data.get('confirm_china_timezone') is True, True)
        basis = _text(data.get('handed_over_basis', ''), '交寄依据', 1000, bool(at))
        if type(data.get('data_complete')) is not bool: issues.append('请明确确认资料完整性')
        if at and checked and at > checked: issues.append('交寄时间晚于资料核对时间')
        if checked and any(e['event_time'] > checked for e in _events(c, identity, True)): issues.append('资料核对时间早于有效事件，请核对')
        if issues: _fail(422, '；'.join(issues))
        s.update(handed_over_at=at, handed_over_basis=basis, checked_until=checked, data_complete=data['data_complete'], version=s['version'] + 1, updated_at=db.now())
        s['history'] = [*s.get('history', []), {'at': db.now(), 'actor': work._actor(user), 'action': '更正运单资料', 'reason': reason, 'before': {k: before.get(k) for k in ['handed_over_at', 'handed_over_basis', 'checked_until', 'data_complete']}, 'after': {k: s.get(k) for k in ['handed_over_at', 'handed_over_basis', 'checked_until', 'data_complete']}}]
        c.execute('UPDATE logistics_shipments SET version=?,data=? WHERE id=?', (s['version'], _json(s), identity)); result = {'shipment': _shipment_view(c, s), **_recheck(c, user, s, 'staff')}
        service.audit(c, user, '更正本地运单资料', identity, before, s)
        return work._remember(c, user, 'logistics:shipment:' + identity, rid, digest, result)
